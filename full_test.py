from ultralytics import YOLO

from flask import Flask, request, jsonify
from werkzeug.serving import make_server

from PIL import Image

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from pathlib import Path

import socket
import struct
import threading
import time
import math
import json

from datetime import datetime


# ============================================================
# 1. 使用者設定區
# ============================================================

# ------------------------------------------------------------
# Robot Network
# ------------------------------------------------------------

ROBOT_IP = "192.168.1.200"
ROBOT_PORT = 5890

SOCKET_TIMEOUT = 5.0


# ------------------------------------------------------------
# 發生 Robot Socket Reset 後
# 等待 TMflow 回到 Listen 的時間
# ------------------------------------------------------------

CONNECTION_RESET_WAIT = 3.0


# ------------------------------------------------------------
# Flask
# ------------------------------------------------------------

HTTP_HOST = "0.0.0.0"
HTTP_PORT = 5000


# ------------------------------------------------------------
# YOLO Model
# ------------------------------------------------------------

MODEL_PATH = Path(
    r"C:\Users\user\Desktop\專題\專題程式\test\best.pt"
)


# ------------------------------------------------------------
# 總測試 Round
#
# 建議：
# 先 10
# 確認正常後再 100
# ------------------------------------------------------------

TOTAL_ROUNDS = 50


# ------------------------------------------------------------
# Robot Move
# ------------------------------------------------------------

SPEED = 80

ACC = 200

BLEND = 0

PRECISE = "false"


# ------------------------------------------------------------
# Robot 到位判定
# ------------------------------------------------------------

POSITION_TOLERANCE_MM = 1.0

MOVE_TIMEOUT = 30.0


# ============================================================
# 重要更新
#
# 原本：
# POLL_INTERVAL = 0.15
#
# 現在：
# 0.30
#
# 降低長時間測試對 Listen / TMSCT 的壓力。
# ============================================================

POLL_INTERVAL = 0.30


STABLE_DELTA_MM = 0.10

STABLE_COUNT_REQUIRED = 3

SETTLE_TIME = 0.30


# ------------------------------------------------------------
# Console 顯示
#
# Robot仍每0.30秒讀一次Pose，
# 但畫面每1秒左右顯示一次。
# ------------------------------------------------------------

POSE_PRINT_INTERVAL = 1.0

SHOW_POSE_PACKET_DEBUG = False


# ------------------------------------------------------------
# Vision
# ------------------------------------------------------------

VISION_TIMEOUT = 60.0


# ------------------------------------------------------------
# YOLO
# ------------------------------------------------------------

YOLO_CONFIDENCE = 0.25


# ------------------------------------------------------------
# 棋盤
# ------------------------------------------------------------

BOARD_ROWS = 10
BOARD_COLS = 9

CURRENT_TURN = "red"


# ============================================================
# 2. Robot 四個正式絕對座標
# ============================================================

P1 = (
    457.60,
    -147.55,
    423.12,
    -177.82,
    2.08,
    60.59,
)


P2 = (
    398.08,
    96.46,
    499.77,
    176.00,
    -0.46,
    88.67,
)


P3 = (
    224.04,
    -214.85,
    534.60,
    -178.61,
    -0.73,
    25.26,
)


P4 = (
    235.63,
    50.15,
    576.70,
    172.35,
    1.06,
    74.35,
)


# ============================================================
# 3. V1 Calibration
#
# 人工90點
# Reference：2048 × 1536
# ============================================================

V1_CALIBRATION_WIDTH = 2048

V1_CALIBRATION_HEIGHT = 1536


V1_CALIBRATION_COL_X = {

    0: 807.5,

    1: 944.0,

    2: 1081.0,

    3: 1218.0,

    4: 1355.5,

    5: 1492.0,

    6: 1629.0,

    7: 1764.5,

    8: 1901.0,

}


V1_CALIBRATION_ROW_Y = {

    0: 155.0,

    1: 295.5,

    2: 433.5,

    3: 571.5,

    4: 709.5,

    5: 832.0,

    6: 969.0,

    7: 1104.0,

    8: 1242.0,

    9: 1380.0,

}


V1_MAX_GRID_DISTANCE = 105.0


# ============================================================
# 4. V2 最新正式 Calibration
#
# Reference：2592 × 1944
# ============================================================

V2_CALIBRATION_WIDTH = 2592

V2_CALIBRATION_HEIGHT = 1944


V2_CALIBRATION_COL_X = {

    0: 1252.0,

    1: 1506.0,

    2: 1765.0,

    3: 2025.0,

}


V2_CALIBRATION_ROW_Y = {

    0: 929.0,

    1: 1189.0,

    2: 1430.0,

    3: 1690.0,

}


V2_MAX_DX = 125.0

V2_MAX_DY = 125.0


# ============================================================
# 5. V2 → Main Board
#
# V2接管：
#
# Main R1~R4
# Main C0~C3
# ============================================================

V2_TO_MAIN_BOARD = {

    (0, 0): (1, 0),
    (0, 1): (1, 1),
    (0, 2): (1, 2),
    (0, 3): (1, 3),

    (1, 0): (2, 0),
    (1, 1): (2, 1),
    (1, 2): (2, 2),
    (1, 3): (2, 3),

    (2, 0): (3, 0),
    (2, 1): (3, 1),
    (2, 2): (3, 2),
    (2, 3): (3, 3),

    (3, 0): (4, 0),
    (3, 1): (4, 1),
    (3, 2): (4, 2),
    (3, 3): (4, 3),

}


V2_MAIN_CELLS = set(
    V2_TO_MAIN_BOARD.values()
)


# ============================================================
# 6. FEN Mapping
# ============================================================

LABEL_TO_FEN = {

    # Black
    "chess-black-king": "k",

    "chess-black-guard": "a",

    "chess-black-elephant": "b",

    "chess-black-horse": "n",

    "chess-black-car": "r",

    "chess-black-cannon": "c",

    "chess-black-soldier": "p",


    # Red
    "chess-red-king": "K",

    "chess-red-guard": "A",

    "chess-red-elephant": "B",

    "chess-red-horse": "N",

    "chess-red-car": "R",

    "chess-red-cannon": "C",

    "chess-red-soldier": "P",

}


# ============================================================
# 7. 測試資料夾
# ============================================================

PROGRAM_DIR = Path(
    __file__
).resolve().parent


RUN_TIMESTAMP = datetime.now().strftime(
    "%Y%m%d_%H%M%S"
)


RUN_DIR = (

    PROGRAM_DIR

    / "TM_4Point_DualVision_YOLO"

    / f"test_{RUN_TIMESTAMP}"

)


IMAGE_DIR = (

    RUN_DIR

    / "images"

)


RUN_DIR.mkdir(
    parents=True,
    exist_ok=True
)


IMAGE_DIR.mkdir(
    parents=True,
    exist_ok=True
)


EXCEL_PATH = (

    RUN_DIR

    / "TM_Robot_YOLO_100Rounds.xlsx"

)


COMMUNICATION_LOG = (

    RUN_DIR

    / "communication.log"

)


# ============================================================
# 8. Flask / Vision State
# ============================================================

app = Flask(
    __name__
)


vision_lock = threading.Lock()


vision1_event = threading.Event()

vision2_event = threading.Event()


vision_state = {

    "round": 0,

    "photo_number": 0,

    "v1_path": None,

    "v2_path": None,

}


# ============================================================
# 9. Communication Log
# ============================================================

log_lock = threading.Lock()


def write_communication_log(
    direction,
    message
):

    timestamp = datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S.%f"
    )


    if isinstance(
        message,
        bytes
    ):

        display_message = repr(
            message
        )

    else:

        display_message = str(
            message
        )


    with log_lock:

        with open(

            COMMUNICATION_LOG,

            "a",

            encoding="utf-8"

        ) as f:

            f.write(

                f"[{timestamp}] "

                f"{direction} "

                f"{display_message}\n"

            )


# ============================================================
# 10. Script ID
# ============================================================

script_counter_lock = threading.Lock()


script_counter = 0


def get_script_id(
    prefix="CMD"
):

    global script_counter


    with script_counter_lock:

        script_counter += 1


        return (

            f"{prefix}"

            f"{script_counter:05d}"

        )


# ============================================================
# 11. 正確 TMSCT Packet
# ============================================================

def make_tmsct_packet(
    script_id,
    script
):

    # --------------------------------------------------------
    # Length只計：
    #
    # ScriptID,Script
    # --------------------------------------------------------

    payload = (

        f"{script_id},"

        f"{script}"

    ).encode(
        "utf-8"
    )


    body = (

        b"TMSCT,"

        +

        str(
            len(payload)
        ).encode(
            "ascii"
        )

        +

        b","

        +

        payload

        +

        b","

    )


    checksum = 0


    for byte in body:

        checksum ^= byte


    packet = (

        b"$"

        +

        body

        +

        b"*"

        +

        f"{checksum:02X}".encode(
            "ascii"
        )

        +

        b"\r\n"

    )


    return packet


# ============================================================
# 12. Connect Robot
# ============================================================

def connect_robot():

    sock = socket.socket(

        socket.AF_INET,

        socket.SOCK_STREAM

    )


    sock.settimeout(
        SOCKET_TIMEOUT
    )


    print(

        f"[Robot] Connecting "

        f"{ROBOT_IP}:"

        f"{ROBOT_PORT}"

    )


    sock.connect(

        (
            ROBOT_IP,
            ROBOT_PORT
        )

    )


    print(
        "[Robot] Connected"
    )


    return sock


# ============================================================
# 13. 判斷 WinError 10054
# ============================================================

def is_connection_reset_error(
    error
):

    # --------------------------------------------------------
    # Python ConnectionResetError
    # --------------------------------------------------------

    if isinstance(
        error,
        ConnectionResetError
    ):

        return True


    # --------------------------------------------------------
    # Windows OSError
    # --------------------------------------------------------

    if isinstance(
        error,
        OSError
    ):

        winerror = getattr(
            error,
            "winerror",
            None
        )


        if winerror == 10054:

            return True


    return False


# ============================================================
# 14. Send Script
# ============================================================

def send_script(
    sock,
    script,
    wait_reply=True,
    prefix="CMD"
):

    script_id = (
        get_script_id(
            prefix
        )
    )


    packet = (
        make_tmsct_packet(

            script_id,

            script

        )
    )


    write_communication_log(

        "SEND",

        packet

    )


    sock.sendall(
        packet
    )


    if not wait_reply:

        return None


    deadline = (

        time.time()

        +

        SOCKET_TIMEOUT

    )


    buffer = b""


    expected_ok = (

        f"{script_id},OK"

    ).encode(
        "ascii"
    )


    while (
        time.time()
        <
        deadline
    ):

        try:

            data = sock.recv(
                4096
            )


        except socket.timeout:

            continue


        if not data:

            continue


        write_communication_log(

            "RECV",

            data

        )


        buffer += data


        # ----------------------------------------------------
        # Controller Error
        # ----------------------------------------------------

        if b"CPERR" in buffer:

            raise RuntimeError(

                f"Robot CPERR："

                f"{buffer!r}"

            )


        # ----------------------------------------------------
        # 確認這一筆 Script ID 的 OK
        # ----------------------------------------------------

        if expected_ok in buffer:

            print(

                f"[Robot OK] "

                f"{script_id} | "

                f"{script}"

            )


            return buffer


    raise TimeoutError(

        f"等待 Robot 回覆逾時："

        f"{script}"

    )


# ============================================================
# 15. DO3
# ============================================================

def set_do3(
    sock,
    state
):

    value = (

        1

        if state

        else 0

    )


    script = (

        f'IO["ControlBox"].'

        f'DO[3] = {value}'

    )


    send_script(

        sock,

        script,

        wait_reply=True,

        prefix="DO3"

    )


    print(

        f"[DO3] "

        f"{'ON' if state else 'OFF'}"

    )


# ============================================================
# 16. Line
# ============================================================

def send_line(
    sock,
    pose
):

    (
        x,
        y,
        z,
        rx,
        ry,
        rz

    ) = pose


    script = (

        f'Line('
        f'"CPP",'
        f'{x:.3f},'
        f'{y:.3f},'
        f'{z:.3f},'
        f'{rx:.3f},'
        f'{ry:.3f},'
        f'{rz:.3f},'
        f'{SPEED},'
        f'{ACC},'
        f'{BLEND},'
        f'{PRECISE}'
        f')'

    )


    send_script(

        sock,

        script,

        wait_reply=True,

        prefix="MOVE"

    )


# ============================================================
# 17. TMSTA Pose Parser
#
# 實際格式：
#
# $TMSTA,27,90,<24 bytes>,*XX\r\n
# ============================================================

def try_parse_pose_packet(
    buffer
):

    marker = b"$TMSTA,"


    search_start = 0


    while True:

        start = buffer.find(

            marker,

            search_start

        )


        if start < 0:

            return None


        # ----------------------------------------------------
        # Length後逗號
        # ----------------------------------------------------

        comma_after_length = buffer.find(

            b",",

            start + len(marker)

        )


        if comma_after_length < 0:

            return None


        # ----------------------------------------------------
        # ID後逗號
        # ----------------------------------------------------

        comma_after_id = buffer.find(

            b",",

            comma_after_length + 1

        )


        if comma_after_id < 0:

            return None


        data_id_bytes = buffer[

            comma_after_length + 1:

            comma_after_id

        ]


        data_id = (

            data_id_bytes

            .decode(
                "ascii",
                errors="ignore"
            )

            .strip()

        )


        # ----------------------------------------------------
        # ListenSend ID = 90
        # ----------------------------------------------------

        if data_id != "90":

            search_start = (

                comma_after_id

                +

                1

            )


            continue


        payload_start = (

            comma_after_id

            +

            1

        )


        # ----------------------------------------------------
        # 6 × float32
        # ----------------------------------------------------

        payload_end = (

            payload_start

            +

            24

        )


        if (

            len(buffer)

            <

            payload_end

        ):

            return None


        pose_bytes = buffer[

            payload_start:

            payload_end

        ]


        try:

            pose = struct.unpack(

                "<6f",

                pose_bytes

            )


        except struct.error:

            return None


        pose = tuple(

            float(value)

            for value
            in pose

        )


        if not all(

            math.isfinite(value)

            for value
            in pose

        ):

            return None


        return pose


# ============================================================
# 18. Read Robot Pose
#
# ★ 本版重要修正
#
# 必須同時收到：
#
# 1. TMSTA Pose
# 2. 此次 POSxxxxx,OK
#
# 才 return。
#
# 不再把 POS OK 留給下一個 command。
# ============================================================

def read_robot_pose(
    sock
):

    script = (

        "ListenSend("

        "90,"

        "Robot[0].CoordRobot"

        ")"

    )


    script_id = (
        get_script_id(
            "POS"
        )
    )


    packet = (
        make_tmsct_packet(

            script_id,

            script

        )
    )


    write_communication_log(

        "SEND",

        packet

    )


    sock.sendall(
        packet
    )


    deadline = (

        time.time()

        +

        SOCKET_TIMEOUT

    )


    buffer = b""


    pose = None


    got_ok = False


    expected_ok = (

        f"{script_id},OK"

    ).encode(
        "ascii"
    )


    while (
        time.time()
        <
        deadline
    ):

        try:

            data = sock.recv(
                4096
            )


        except socket.timeout:

            continue


        if not data:

            continue


        write_communication_log(

            "RECV",

            data

        )


        buffer += data


        # ----------------------------------------------------
        # CPERR
        # ----------------------------------------------------

        if b"CPERR" in buffer:

            raise RuntimeError(

                f"Robot CPERR："

                f"{buffer!r}"

            )


        # ----------------------------------------------------
        # Debug
        # ----------------------------------------------------

        if SHOW_POSE_PACKET_DEBUG:

            if b"$TMSTA" in data:

                print(
                    "[Pose] 收到 TMSTA"
                )


            if expected_ok in buffer:

                print(

                    f"[Pose] 收到 "

                    f"{script_id} OK"

                )


        # ----------------------------------------------------
        # Pose
        # ----------------------------------------------------

        if pose is None:

            parsed_pose = (
                try_parse_pose_packet(
                    buffer
                )
            )


            if parsed_pose is not None:

                pose = (
                    parsed_pose
                )


        # ----------------------------------------------------
        # 當次 POS OK
        # ----------------------------------------------------

        if expected_ok in buffer:

            got_ok = True


        # ====================================================
        # 兩者都收到才 return
        # ====================================================

        if (

            pose is not None

            and

            got_ok

        ):

            return pose


    # ========================================================
    # Timeout Debug
    # ========================================================

    missing = []


    if pose is None:

        missing.append(
            "TMSTA Pose"
        )


    if not got_ok:

        missing.append(
            f"{script_id} OK"
        )


    raise TimeoutError(

        f"無法完整取得 Robot Pose，缺少："

        f"{', '.join(missing)}"

    )


# ============================================================
# 19. XYZ Distance
# ============================================================

def xyz_distance(
    pose_a,
    pose_b
):

    return math.sqrt(

        (
            pose_a[0]
            -
            pose_b[0]
        ) ** 2

        +

        (
            pose_a[1]
            -
            pose_b[1]
        ) ** 2

        +

        (
            pose_a[2]
            -
            pose_b[2]
        ) ** 2

    )


# ============================================================
# 20. Wait Until Pose
# ============================================================

def wait_until_pose(
    sock,
    target_pose,
    point_name
):

    deadline = (

        time.time()

        +

        MOVE_TIMEOUT

    )


    last_pose = None


    stable_count = 0


    best_error = float(
        "inf"
    )


    read_fail_count = 0


    last_print_time = 0.0


    while (
        time.time()
        <
        deadline
    ):

        try:

            current_pose = (
                read_robot_pose(
                    sock
                )
            )


        except TimeoutError as error:

            read_fail_count += 1


            print(

                f"[{point_name}] "

                f"Pose Read Fail "

                f"#{read_fail_count} | "

                f"{error}"

            )


            if (
                read_fail_count
                >=
                3
            ):

                raise TimeoutError(

                    f"{point_name} "

                    f"連續3次無法完整取得 "

                    f"Robot Pose"

                )


            time.sleep(
                0.30
            )


            continue


        # ----------------------------------------------------
        # 成功後歸零
        # ----------------------------------------------------

        read_fail_count = 0


        # ----------------------------------------------------
        # XYZ Error
        # ----------------------------------------------------

        error = xyz_distance(

            current_pose,

            target_pose

        )


        best_error = min(

            best_error,

            error

        )


        # ----------------------------------------------------
        # Delta
        # ----------------------------------------------------

        if last_pose is None:

            delta = float(
                "inf"
            )

        else:

            delta = xyz_distance(

                current_pose,

                last_pose

            )


        # ----------------------------------------------------
        # Stable
        # ----------------------------------------------------

        if (

            error
            <=
            POSITION_TOLERANCE_MM

            and

            delta
            <=
            STABLE_DELTA_MM

        ):

            stable_count += 1

        else:

            stable_count = 0


        # ----------------------------------------------------
        # Console
        # ----------------------------------------------------

        now = time.time()


        if (

            now
            -
            last_print_time

            >=

            POSE_PRINT_INTERVAL

        ):

            print(

                f"[{point_name}] "

                f"XYZ=("
                f"{current_pose[0]:.3f}, "
                f"{current_pose[1]:.3f}, "
                f"{current_pose[2]:.3f}"
                f") "

                f"Error="
                f"{error:.4f} mm "

                f"Delta="
                f"{delta:.4f} mm "

                f"Stable="
                f"{stable_count}/"
                f"{STABLE_COUNT_REQUIRED}"

            )


            last_print_time = (
                now
            )


        # ----------------------------------------------------
        # Reached
        # ----------------------------------------------------

        if (

            stable_count
            >=
            STABLE_COUNT_REQUIRED

        ):

            time.sleep(
                SETTLE_TIME
            )


            final_error = xyz_distance(

                current_pose,

                target_pose

            )


            print(

                f"[{point_name}] "

                f"Reached | "

                f"XYZ=("
                f"{current_pose[0]:.3f}, "
                f"{current_pose[1]:.3f}, "
                f"{current_pose[2]:.3f}"
                f") | "

                f"Error="
                f"{final_error:.4f} mm"

            )


            return (

                current_pose,

                final_error

            )


        last_pose = (
            current_pose
        )


        time.sleep(
            POLL_INTERVAL
        )


    raise TimeoutError(

        f"{point_name} "

        f"移動逾時 | "

        f"Best Error="

        f"{best_error:.4f} mm"

    )


# ============================================================
# 21. Move + Wait
# ============================================================

def move_and_wait(
    sock,
    pose,
    point_name
):

    print()

    print(

        f"--- Move "
        f"{point_name} ---"

    )


    send_line(

        sock,

        pose

    )


    return wait_until_pose(

        sock,

        pose,

        point_name

    )


# ============================================================
# 22. ScriptExit
# ============================================================

def send_script_exit(
    sock
):

    print()

    print(
        "[Robot] ScriptExit()"
    )


    send_script(

        sock,

        "ScriptExit()",

        wait_reply=True,

        prefix="EXIT"

    )


# ============================================================
# 23. Prepare Vision
# ============================================================

def prepare_vision_round(
    round_no
):

    vision1_event.clear()

    vision2_event.clear()


    with vision_lock:

        vision_state[
            "round"
        ] = (
            round_no
        )


        vision_state[
            "photo_number"
        ] = 0


        vision_state[
            "v1_path"
        ] = None


        vision_state[
            "v2_path"
        ] = None


# ============================================================
# 24. Flask Health
# ============================================================

@app.route(
    "/",
    methods=[
        "GET"
    ]
)
def health():

    return jsonify({

        "message":
            "TM Robot Vision Server Running"

    })


# ============================================================
# 25. TMvision Endpoint
# ============================================================

@app.route(
    "/api/vision/tmvision/detect",
    methods=[
        "GET",
        "POST"
    ]
)
def tmvision_detect():

    if (
        request.method
        ==
        "GET"
    ):

        return jsonify({

            "message":
                "TMvision endpoint ready"

        })


    # --------------------------------------------------------
    # file / image
    # --------------------------------------------------------

    uploaded_file = (
        request.files.get(
            "file"
        )
    )


    if uploaded_file is None:

        uploaded_file = (
            request.files.get(
                "image"
            )
        )


    if uploaded_file is None:

        return jsonify({

            "message":
                "image file not found",

            "annotations":
                []

        }), 400


    image_bytes = (
        uploaded_file.read()
    )


    # --------------------------------------------------------
    # Round / Photo
    # --------------------------------------------------------

    with vision_lock:

        round_no = (

            vision_state[
                "round"
            ]

        )


        vision_state[
            "photo_number"
        ] += 1


        photo_no = (

            vision_state[
                "photo_number"
            ]

        )


    # --------------------------------------------------------
    # 一個Round只接受兩張
    # --------------------------------------------------------

    if photo_no > 2:

        print(

            f"[Flask] "

            f"Round {round_no} "

            f"收到多餘影像 "

            f"#{photo_no}"

        )


        return jsonify({

            "message":
                "extra image ignored",

            "annotations":
                []

        }), 200


    view_no = (
        photo_no
    )


    # ========================================================
    # File Name
    #
    # TM_R001_V1_20261008_1_1.jpg
    # ========================================================

    date_text = (
        datetime.now().strftime(
            "%Y%m%d"
        )
    )


    filename = (

        f"TM_R{round_no:03d}_"

        f"V{view_no}_"

        f"{date_text}_"

        f"{round_no}_"

        f"{photo_no}.jpg"

    )


    image_path = (

        IMAGE_DIR

        /

        filename

    )


    with open(

        image_path,

        "wb"

    ) as f:

        f.write(
            image_bytes
        )


    print()

    print(
        "=" * 90
    )


    print(

        f"[Flask] "

        f"Round {round_no} "

        f"Vision {view_no}"

    )


    print(

        f"Saved："

        f"{image_path}"

    )


    print(
        "=" * 90
    )


    # --------------------------------------------------------
    # Size
    # --------------------------------------------------------

    try:

        with Image.open(
            image_path
        ) as image:

            image_w, image_h = (
                image.size
            )


    except Exception:

        image_w = 2592

        image_h = 1944


    # --------------------------------------------------------
    # Event
    # --------------------------------------------------------

    with vision_lock:

        if view_no == 1:

            vision_state[
                "v1_path"
            ] = (
                image_path
            )


            vision1_event.set()


        elif view_no == 2:

            vision_state[
                "v2_path"
            ] = (
                image_path
            )


            vision2_event.set()


    # ========================================================
    # TMvision External Detection
    # ========================================================

    annotation = {

        "label":
            "capture_ok",

        "score":
            0.99,

        "box_cx":
            image_w / 2,

        "box_cy":
            image_h / 2,

        "box_w":
            100,

        "box_h":
            100,

        "rotation":
            0.0,

    }


    response_json = {

        "message":
            "success",

        "image_w":
            image_w,

        "image_h":
            image_h,

        "annotations":
            [
                annotation
            ],

    }


    return jsonify(
        response_json
    ), 200


# ============================================================
# 26. Flask Thread
# ============================================================

class FlaskServerThread(
    threading.Thread
):

    def __init__(
        self
    ):

        super().__init__(
            daemon=True
        )


        self.server = make_server(

            HTTP_HOST,

            HTTP_PORT,

            app,

            threaded=True

        )


    def run(
        self
    ):

        print(

            f"[Flask] Running："

            f"http://"

            f"{HTTP_HOST}:"

            f"{HTTP_PORT}"

        )


        self.server.serve_forever()


    def shutdown(
        self
    ):

        self.server.shutdown()


# ============================================================
# 27. Wait V1 + V2
# ============================================================

def wait_for_two_images(
    round_no
):

    print()

    print(

        f"[Vision] "

        f"Waiting Round "

        f"{round_no} "

        f"V1 + V2..."

    )


    deadline = (

        time.time()

        +

        VISION_TIMEOUT

    )


    while (
        time.time()
        <
        deadline
    ):

        if (

            vision1_event.is_set()

            and

            vision2_event.is_set()

        ):

            with vision_lock:

                v1_path = (

                    vision_state[
                        "v1_path"
                    ]

                )


                v2_path = (

                    vision_state[
                        "v2_path"
                    ]

                )


            if (

                v1_path is not None

                and

                v2_path is not None

            ):

                print(

                    f"[Vision] "

                    f"Round {round_no} "

                    f"2/2 Images Ready"

                )


                return (

                    Path(
                        v1_path
                    ),

                    Path(
                        v2_path
                    )

                )


        time.sleep(
            0.10
        )


    raise TimeoutError(

        f"Round {round_no} "

        f"等待 V1 / V2 "

        f"圖片逾時"

    )


# ============================================================
# 28. Empty Board
# ============================================================

def create_empty_board():

    return [

        [

            None

            for _ in range(
                BOARD_COLS
            )

        ]

        for _ in range(
            BOARD_ROWS
        )

    ]


# ============================================================
# 29. Build V1 Grid
# ============================================================

def build_v1_grid(
    width,
    height
):

    sx = (

        width

        /

        V1_CALIBRATION_WIDTH

    )


    sy = (

        height

        /

        V1_CALIBRATION_HEIGHT

    )


    grid = {}


    for row in range(
        BOARD_ROWS
    ):

        for col in range(
            BOARD_COLS
        ):

            x = (

                V1_CALIBRATION_COL_X[
                    col
                ]

                *

                sx

            )


            y = (

                V1_CALIBRATION_ROW_Y[
                    row
                ]

                *

                sy

            )


            grid[
                (row, col)
            ] = (

                x,
                y

            )


    return grid


# ============================================================
# 30. Build V2 Grid
# ============================================================

def build_v2_grid(
    width,
    height
):

    sx = (

        width

        /

        V2_CALIBRATION_WIDTH

    )


    sy = (

        height

        /

        V2_CALIBRATION_HEIGHT

    )


    grid = {}


    for row in range(
        4
    ):

        for col in range(
            4
        ):

            x = (

                V2_CALIBRATION_COL_X[
                    col
                ]

                *

                sx

            )


            y = (

                V2_CALIBRATION_ROW_Y[
                    row
                ]

                *

                sy

            )


            grid[
                (row, col)
            ] = (

                x,
                y

            )


    return grid


# ============================================================
# 31. Find Nearest Grid
# ============================================================

def find_nearest_grid_cell(
    cx,
    cy,
    grid
):

    best_cell = None


    best_distance = float(
        "inf"
    )


    for cell, (
        gx,
        gy
    ) in grid.items():

        distance = math.hypot(

            cx - gx,

            cy - gy

        )


        if (

            distance

            <

            best_distance

        ):

            best_cell = (
                cell
            )


            best_distance = (
                distance
            )


    return (

        best_cell,

        best_distance

    )


# ============================================================
# 32. YOLO
# ============================================================

def run_yolo(
    model,
    image_path
):

    # --------------------------------------------------------
    # 原圖名稱 + _YOLO
    # --------------------------------------------------------

    yolo_path = image_path.with_name(

        f"{image_path.stem}"

        f"_YOLO"

        f"{image_path.suffix}"

    )


    results = model.predict(

        source=str(
            image_path
        ),

        conf=YOLO_CONFIDENCE,

        verbose=False

    )


    result = (
        results[0]
    )


    result.save(

        filename=str(
            yolo_path
        )

    )


    detections = []


    if (

        result.boxes is None

        or

        len(
            result.boxes
        ) == 0

    ):

        return (

            detections,

            yolo_path

        )


    for index, box in enumerate(
        result.boxes
    ):

        cls_id = int(
            box.cls[0]
        )


        label = (
            model.names[
                cls_id
            ]
        )


        confidence = float(
            box.conf[0]
        )


        x1, y1, x2, y2 = (

            box.xyxy[0]

            .cpu()

            .numpy()

            .tolist()

        )


        cx = (
            x1 + x2
        ) / 2


        cy = (
            y1 + y2
        ) / 2


        detections.append({

            "index":
                index,

            "label":
                label,

            "confidence":
                confidence,

            "cx":
                cx,

            "cy":
                cy,

        })


    return (

        detections,

        yolo_path

    )


# ============================================================
# 33. V1 Mapping
# ============================================================

def map_v1(
    detections,
    grid
):

    board = (
        create_empty_board()
    )


    cell_info = {}


    stats = {

        "detection":
            len(
                detections
            ),

        "mapped":
            0,

        "reserved_v2":
            0,

        "ignored":
            0,

        "collision":
            0,

    }


    for det in detections:

        (
            cell,
            distance

        ) = find_nearest_grid_cell(

            det[
                "cx"
            ],

            det[
                "cy"
            ],

            grid

        )


        row, col = (
            cell
        )


        if (

            distance

            >

            V1_MAX_GRID_DISTANCE

        ):

            stats[
                "ignored"
            ] += 1


            continue


        # ----------------------------------------------------
        # V2區域不使用V1
        # ----------------------------------------------------

        if (
            cell
            in
            V2_MAIN_CELLS
        ):

            stats[
                "reserved_v2"
            ] += 1


            continue


        # ----------------------------------------------------
        # Collision
        # ----------------------------------------------------

        if (
            cell
            in
            cell_info
        ):

            stats[
                "collision"
            ] += 1


            old = (
                cell_info[
                    cell
                ]
            )


            if (

                det[
                    "confidence"
                ]

                <=

                old[
                    "confidence"
                ]

            ):

                continue


        board[
            row
        ][
            col
        ] = (

            det[
                "label"
            ]

        )


        cell_info[
            cell
        ] = (
            det
        )


    stats[
        "mapped"
    ] = len(
        cell_info
    )


    return (

        board,

        cell_info,

        stats

    )


# ============================================================
# 34. V2 Mapping
# ============================================================

def map_v2(
    detections,
    grid
):

    cell_info = {}


    stats = {

        "detection":
            len(
                detections
            ),

        "mapped":
            0,

        "ignored":
            0,

        "collision":
            0,

    }


    for det in detections:

        (
            local_cell,
            distance

        ) = find_nearest_grid_cell(

            det[
                "cx"
            ],

            det[
                "cy"
            ],

            grid

        )


        gx, gy = (
            grid[
                local_cell
            ]
        )


        dx = abs(

            det[
                "cx"
            ]

            -

            gx

        )


        dy = abs(

            det[
                "cy"
            ]

            -

            gy

        )


        if (

            dx
            >
            V2_MAX_DX

            or

            dy
            >
            V2_MAX_DY

        ):

            stats[
                "ignored"
            ] += 1


            continue


        main_cell = (

            V2_TO_MAIN_BOARD[
                local_cell
            ]

        )


        if (
            main_cell
            in
            cell_info
        ):

            stats[
                "collision"
            ] += 1


            old = (
                cell_info[
                    main_cell
                ]
            )


            if (

                det[
                    "confidence"
                ]

                <=

                old[
                    "confidence"
                ]

            ):

                continue


        cell_info[
            main_cell
        ] = det


    stats[
        "mapped"
    ] = len(
        cell_info
    )


    return (

        cell_info,

        stats

    )


# ============================================================
# 35. Fusion
# ============================================================

def fuse_boards(
    v1_board,
    v2_info
):

    final_board = [

        row.copy()

        for row in v1_board

    ]


    # --------------------------------------------------------
    # V2接管區先清空
    # --------------------------------------------------------

    for (
        row,
        col
    ) in V2_MAIN_CELLS:

        final_board[
            row
        ][
            col
        ] = None


    # --------------------------------------------------------
    # V2填入
    # --------------------------------------------------------

    for (
        row,
        col
    ), det in v2_info.items():

        final_board[
            row
        ][
            col
        ] = (

            det[
                "label"
            ]

        )


    return final_board


# ============================================================
# 36. Board → FEN
# ============================================================

def board_to_fen(
    board
):

    fen_rows = []


    for board_row in board:

        fen_row = ""


        empty = 0


        for piece in board_row:

            if piece is None:

                empty += 1


                continue


            if empty > 0:

                fen_row += str(
                    empty
                )


                empty = 0


            fen_row += LABEL_TO_FEN.get(
                piece,
                "?"
            )


        if empty > 0:

            fen_row += str(
                empty
            )


        fen_rows.append(
            fen_row
        )


    side = (

        "w"

        if
        CURRENT_TURN
        ==
        "red"

        else
        "b"

    )


    return (

        "/".join(
            fen_rows
        )

        +

        f" {side} - - 0 1"

    )


# ============================================================
# 37. Board JSON
# ============================================================

def board_to_json(
    board
):

    matrix = []


    for row in board:

        matrix_row = []


        for piece in row:

            if piece is None:

                matrix_row.append(
                    "."
                )

            else:

                matrix_row.append(

                    piece.replace(
                        "chess-",
                        ""
                    )

                )


        matrix.append(
            matrix_row
        )


    return json.dumps(

        matrix,

        ensure_ascii=False

    )


# ============================================================
# 38. Count Pieces
# ============================================================

def count_board_pieces(
    board
):

    return sum(

        1

        for row in board

        for piece in row

        if piece is not None

    )


# ============================================================
# 39. Dual View YOLO
# ============================================================

def process_dual_view_yolo(
    model,
    v1_path,
    v2_path,
    round_no
):

    print()

    print(
        "=" * 110
    )


    print(

        f"ROUND "
        f"{round_no:03d} "

        f"YOLO START"

    )


    print(
        "=" * 110
    )


    start = (
        time.perf_counter()
    )


    # --------------------------------------------------------
    # Image Size
    # --------------------------------------------------------

    with Image.open(
        v1_path
    ) as image:

        v1_width, v1_height = (
            image.size
        )


    with Image.open(
        v2_path
    ) as image:

        v2_width, v2_height = (
            image.size
        )


    # --------------------------------------------------------
    # Grid
    # --------------------------------------------------------

    v1_grid = build_v1_grid(

        v1_width,

        v1_height

    )


    v2_grid = build_v2_grid(

        v2_width,

        v2_height

    )


    # --------------------------------------------------------
    # YOLO V1
    # --------------------------------------------------------

    (
        v1_detections,
        v1_yolo_path

    ) = run_yolo(

        model,

        v1_path

    )


    print(

        f"[YOLO V1] "

        f"Detection="

        f"{len(v1_detections)}"

    )


    # --------------------------------------------------------
    # YOLO V2
    # --------------------------------------------------------

    (
        v2_detections,
        v2_yolo_path

    ) = run_yolo(

        model,

        v2_path

    )


    print(

        f"[YOLO V2] "

        f"Detection="

        f"{len(v2_detections)}"

    )


    # --------------------------------------------------------
    # Mapping
    # --------------------------------------------------------

    (
        v1_board,
        v1_info,
        v1_stats

    ) = map_v1(

        v1_detections,

        v1_grid

    )


    (
        v2_info,
        v2_stats

    ) = map_v2(

        v2_detections,

        v2_grid

    )


    # --------------------------------------------------------
    # Fusion
    # --------------------------------------------------------

    final_board = fuse_boards(

        v1_board,

        v2_info

    )


    # --------------------------------------------------------
    # FEN / UCCI
    # --------------------------------------------------------

    fen = board_to_fen(
        final_board
    )


    ucci = (

        f"position fen "

        f"{fen}"

    )


    final_piece_count = (
        count_board_pieces(
            final_board
        )
    )


    yolo_time = (

        time.perf_counter()

        -

        start

    )


    # --------------------------------------------------------
    # Print Matrix
    # --------------------------------------------------------

    print()

    print(
        "FINAL MATRIX"
    )


    print(
        "-" * 110
    )


    for row_index, row in enumerate(
        final_board
    ):

        row_text = []


        for piece in row:

            if piece is None:

                row_text.append(
                    "."
                )

            else:

                row_text.append(

                    piece.replace(
                        "chess-",
                        ""
                    )

                )


        print(

            f"R{row_index}: "

            f"{row_text}"

        )


    print()

    print(
        f"FEN：{fen}"
    )


    print()

    print(
        f"UCCI：{ucci}"
    )


    print()

    print(

        f"YOLO Time："

        f"{yolo_time:.3f} sec"

    )


    return {

        "v1_yolo_path":
            v1_yolo_path,

        "v2_yolo_path":
            v2_yolo_path,

        "v1_stats":
            v1_stats,

        "v2_stats":
            v2_stats,

        "final_board":
            final_board,

        "final_piece_count":
            final_piece_count,

        "fen":
            fen,

        "ucci":
            ucci,

        "yolo_time":
            yolo_time,

    }


# ============================================================
# 40. Excel
# ============================================================

def create_excel():

    wb = Workbook()


    ws = wb.active


    ws.title = (
        "100_Rounds"
    )


    headers = [

        "Round",

        "Start_Time",

        "End_Time",

        "Round_Time_sec",

        "V1_Image",

        "V2_Image",

        "V1_YOLO_Image",

        "V2_YOLO_Image",

        "V1_Detection",

        "V1_Mapped",

        "V1_Reserved_V2",

        "V1_Ignore",

        "V1_Collision",

        "V2_Detection",

        "V2_Mapped",

        "V2_Ignore",

        "V2_Collision",

        "Final_Piece_Count",

        "FEN",

        "UCCI",

        "Final_Matrix",

        "YOLO_Time_sec",

        "P1_Error_mm",

        "P2_Error_mm",

        "P3_Error_mm",

        "P4_Error_mm",

        "Result",

        "Error",

    ]


    ws.append(
        headers
    )


    for cell in ws[1]:

        cell.font = Font(
            bold=True
        )


        cell.alignment = Alignment(

            horizontal="center",

            vertical="center"

        )


    for index in range(
        1,
        len(headers) + 1
    ):

        ws.column_dimensions[
            get_column_letter(
                index
            )
        ].width = 20


    ws.column_dimensions["E"].width = 45

    ws.column_dimensions["F"].width = 45

    ws.column_dimensions["G"].width = 50

    ws.column_dimensions["H"].width = 50

    ws.column_dimensions["S"].width = 80

    ws.column_dimensions["T"].width = 90

    ws.column_dimensions["U"].width = 120

    ws.column_dimensions["AB"].width = 60


    wb.save(
        EXCEL_PATH
    )


    return (

        wb,

        ws

    )


# ============================================================
# 41. Append Excel
# ============================================================

def append_excel_row(
    wb,
    ws,
    round_no,
    start_time,
    end_time,
    cycle_time,
    v1_path,
    v2_path,
    yolo_result,
    point_errors,
    result_text,
    error_text
):

    if yolo_result is None:

        v1_yolo = ""

        v2_yolo = ""


        v1_stats = {

            "detection": "",

            "mapped": "",

            "reserved_v2": "",

            "ignored": "",

            "collision": "",

        }


        v2_stats = {

            "detection": "",

            "mapped": "",

            "ignored": "",

            "collision": "",

        }


        final_piece_count = ""

        fen = ""

        ucci = ""

        matrix = ""

        yolo_time = ""


    else:

        v1_yolo = (

            yolo_result[
                "v1_yolo_path"
            ].name

        )


        v2_yolo = (

            yolo_result[
                "v2_yolo_path"
            ].name

        )


        v1_stats = (

            yolo_result[
                "v1_stats"
            ]

        )


        v2_stats = (

            yolo_result[
                "v2_stats"
            ]

        )


        final_piece_count = (

            yolo_result[
                "final_piece_count"
            ]

        )


        fen = (

            yolo_result[
                "fen"
            ]

        )


        ucci = (

            yolo_result[
                "ucci"
            ]

        )


        matrix = board_to_json(

            yolo_result[
                "final_board"
            ]

        )


        yolo_time = round(

            yolo_result[
                "yolo_time"
            ],

            4

        )


    ws.append([

        round_no,

        start_time,

        end_time,

        round(
            cycle_time,
            4
        ),

        (
            v1_path.name
            if v1_path
            else ""
        ),

        (
            v2_path.name
            if v2_path
            else ""
        ),

        v1_yolo,

        v2_yolo,

        v1_stats[
            "detection"
        ],

        v1_stats[
            "mapped"
        ],

        v1_stats[
            "reserved_v2"
        ],

        v1_stats[
            "ignored"
        ],

        v1_stats[
            "collision"
        ],

        v2_stats[
            "detection"
        ],

        v2_stats[
            "mapped"
        ],

        v2_stats[
            "ignored"
        ],

        v2_stats[
            "collision"
        ],

        final_piece_count,

        fen,

        ucci,

        matrix,

        yolo_time,

        point_errors.get(
            "P1",
            ""
        ),

        point_errors.get(
            "P2",
            ""
        ),

        point_errors.get(
            "P3",
            ""
        ),

        point_errors.get(
            "P4",
            ""
        ),

        result_text,

        error_text,

    ])


    # --------------------------------------------------------
    # 每Round都存一次
    # --------------------------------------------------------

    wb.save(
        EXCEL_PATH
    )


# ============================================================
# 42. Robot Round
# ============================================================

def run_robot_round(
    round_no
):

    sock = None


    point_errors = {}


    try:

        # ----------------------------------------------------
        # Connect
        # ----------------------------------------------------

        sock = (
            connect_robot()
        )


        # ----------------------------------------------------
        # Start：
        # DO3 OFF
        # ----------------------------------------------------

        set_do3(

            sock,

            False

        )


        # ----------------------------------------------------
        # P1
        # ----------------------------------------------------

        _, error = move_and_wait(

            sock,

            P1,

            "P1"

        )


        point_errors[
            "P1"
        ] = round(
            error,
            4
        )


        # ----------------------------------------------------
        # P2
        # ----------------------------------------------------

        _, error = move_and_wait(

            sock,

            P2,

            "P2"

        )


        point_errors[
            "P2"
        ] = round(
            error,
            4
        )


        # ----------------------------------------------------
        # P2 → DO3 ON
        # ----------------------------------------------------

        set_do3(

            sock,

            True

        )


        # ----------------------------------------------------
        # P3
        # ----------------------------------------------------

        _, error = move_and_wait(

            sock,

            P3,

            "P3"

        )


        point_errors[
            "P3"
        ] = round(
            error,
            4
        )


        # ----------------------------------------------------
        # P3 → DO3 OFF
        # ----------------------------------------------------

        set_do3(

            sock,

            False

        )


        # ----------------------------------------------------
        # P4
        # ----------------------------------------------------

        _, error = move_and_wait(

            sock,

            P4,

            "P4"

        )


        point_errors[
            "P4"
        ] = round(
            error,
            4
        )


        # ----------------------------------------------------
        # Safety DO3 OFF
        # ----------------------------------------------------

        set_do3(

            sock,

            False

        )


        # ----------------------------------------------------
        # 一個Round只送一次
        # ----------------------------------------------------

        send_script_exit(
            sock
        )


        return (
            point_errors
        )


    except Exception:

        # ----------------------------------------------------
        # 發生錯誤時：
        # 若Socket仍活著嘗試DO3 OFF
        # ----------------------------------------------------

        if sock is not None:

            try:

                set_do3(

                    sock,

                    False

                )

            except Exception:

                pass


        raise


    finally:

        # ----------------------------------------------------
        # Close Socket
        # ----------------------------------------------------

        if sock is not None:

            try:

                sock.shutdown(
                    socket.SHUT_RDWR
                )

            except Exception:

                pass


            try:

                sock.close()

            except Exception:

                pass


# ============================================================
# 43. Main
# ============================================================

def main():

    print()

    print(
        "=" * 120
    )


    print(
        "TM Robot + Dual Vision + YOLO"
    )


    print(

        f"TOTAL ROUNDS = "

        f"{TOTAL_ROUNDS}"

    )


    print(
        "=" * 120
    )


    # ========================================================
    # Model Check
    # ========================================================

    if not MODEL_PATH.exists():

        raise FileNotFoundError(

            f"找不到 YOLO Model：\n"

            f"{MODEL_PATH}"

        )


    # ========================================================
    # Flask
    # ========================================================

    flask_server = (
        FlaskServerThread()
    )


    flask_server.start()


    time.sleep(
        1.0
    )


    # ========================================================
    # YOLO只載入一次
    # ========================================================

    print()

    print(
        "[YOLO] Loading Model..."
    )


    model = YOLO(

        str(
            MODEL_PATH
        )

    )


    print(
        "[YOLO] Model Loaded"
    )


    # ========================================================
    # Excel
    # ========================================================

    (
        wb,
        ws

    ) = create_excel()


    passed_rounds = 0

    failed_rounds = 0


    total_test_start = (
        time.perf_counter()
    )


    # ========================================================
    # Round Loop
    # ========================================================

    for round_no in range(

        1,

        TOTAL_ROUNDS + 1

    ):

        print()

        print()


        print(
            "#" * 120
        )


        print(

            f"ROUND "

            f"{round_no:03d}"

            f"/"

            f"{TOTAL_ROUNDS:03d}"

        )


        print(
            "#" * 120
        )


        round_start_perf = (
            time.perf_counter()
        )


        start_time_text = (

            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )

        )


        v1_path = None

        v2_path = None


        yolo_result = None


        point_errors = {}


        result_text = (
            "FAIL"
        )


        error_text = ""


        connection_reset_happened = False


        try:

            # =================================================
            # Vision準備
            # =================================================

            prepare_vision_round(
                round_no
            )


            # =================================================
            # Robot
            # =================================================

            point_errors = (
                run_robot_round(
                    round_no
                )
            )


            # =================================================
            # 等待兩張照片
            # =================================================

            (
                v1_path,
                v2_path

            ) = wait_for_two_images(
                round_no
            )


            print()

            print(
                f"V1："
                f"{v1_path.name}"
            )


            print(
                f"V2："
                f"{v2_path.name}"
            )


            # =================================================
            # YOLO
            # =================================================

            yolo_result = (
                process_dual_view_yolo(

                    model,

                    v1_path,

                    v2_path,

                    round_no

                )
            )


            result_text = (
                "PASS"
            )


            passed_rounds += 1


        except Exception as error:

            error_text = (
                str(error)
            )


            failed_rounds += 1


            connection_reset_happened = (
                is_connection_reset_error(
                    error
                )
            )


            print()

            print(
                "!" * 70
            )


            print(

                f"[ROUND "

                f"{round_no:03d}] "

                f"FAIL"

            )


            print(
                error_text
            )


            # =================================================
            # Connection Reset Recovery
            # =================================================

            if connection_reset_happened:

                print()

                print(

                    "[Recovery] "

                    "偵測到 WinError 10054 / "

                    "ConnectionResetError"

                )


                print(

                    "[Recovery] "

                    f"等待 "

                    f"{CONNECTION_RESET_WAIT:.1f} 秒，"

                    "讓 TMflow 回到 Listen..."

                )


            print(
                "!" * 70
            )


        # ====================================================
        # Round End
        # ====================================================

        end_time_text = (

            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )

        )


        cycle_time = (

            time.perf_counter()

            -

            round_start_perf

        )


        # ====================================================
        # Excel
        # ====================================================

        append_excel_row(

            wb=wb,

            ws=ws,

            round_no=round_no,

            start_time=start_time_text,

            end_time=end_time_text,

            cycle_time=cycle_time,

            v1_path=v1_path,

            v2_path=v2_path,

            yolo_result=yolo_result,

            point_errors=point_errors,

            result_text=result_text,

            error_text=error_text,

        )


        # ====================================================
        # Summary
        # ====================================================

        print()

        print(
            "-" * 100
        )


        print(

            f"Round "

            f"{round_no:03d} "

            f"{result_text}"

        )


        print(

            f"Round Time："

            f"{cycle_time:.3f} sec"

        )


        if yolo_result is not None:

            print()

            print(

                f"FEN："

                f"{yolo_result['fen']}"

            )


            print(

                f"UCCI："

                f"{yolo_result['ucci']}"

            )


        print()

        print(

            f"PASS："

            f"{passed_rounds}"

        )


        print(

            f"FAIL："

            f"{failed_rounds}"

        )


        print(
            "-" * 100
        )


        # ====================================================
        # ★ Connection Reset Recovery
        #
        # 必須放在 Excel 已存完之後，
        # 下一Round開始之前。
        # ====================================================

        if connection_reset_happened:

            time.sleep(
                CONNECTION_RESET_WAIT
            )


    # ========================================================
    # Finished
    # ========================================================

    total_test_time = (

        time.perf_counter()

        -

        total_test_start

    )


    wb.save(
        EXCEL_PATH
    )


    print()

    print()


    print(
        "=" * 120
    )


    print(
        "TEST FINISHED"
    )


    print(
        "=" * 120
    )


    print(

        f"PASS："

        f"{passed_rounds}"

        f"/"

        f"{TOTAL_ROUNDS}"

    )


    print(

        f"FAIL："

        f"{failed_rounds}"

        f"/"

        f"{TOTAL_ROUNDS}"

    )


    print(

        f"Total Time："

        f"{total_test_time:.3f} sec"

    )


    print()

    print(

        f"Image Folder：\n"

        f"{IMAGE_DIR}"

    )


    print()

    print(

        f"Excel：\n"

        f"{EXCEL_PATH}"

    )


    print()

    print(

        f"Communication Log：\n"

        f"{COMMUNICATION_LOG}"

    )


    print(
        "=" * 120
    )


# ============================================================
# 44. Start
# ============================================================

if __name__ == "__main__":

    main()