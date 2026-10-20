"""Rebuild derived Excel reports without modifying historical evidence."""
from __future__ import annotations
import argparse
import copy
import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import Workbook
from tmchess.board import board_to_fen
from tmchess.config import CONFIG
from tmchess.record_audit import audit_history
from tmchess.records import GameRecordWriter, _datetime_value


def enrich(source, directory, *, correct_metadata=False):
    p = copy.deepcopy(source)
    frames, issues = audit_history(source)
    corrections = []
    def change(obj, key, value, location):
        old = obj.get(key)
        if old != value:
            corrections.append(dict(location=location, field=key, original=old, derived=value))
            obj[key] = value
    # A complete replay must reconstruct the stored board before any metadata
    # is derived. Do not guess an initial side or fill missing moves.
    if correct_metadata and frames and len(frames) == len(p.get('move_history', [])) + 1 and frames[-1]['board'] == p.get('board'):
        for i, m in enumerate(p['move_history']):
            change(m, 'side', frames[i]['side'], f'move:{i+1}')
            change(m, 'fen_before', board_to_fen(frames[i]['board'],frames[i]['side']), f'move:{i+1}')
            change(m, 'fen_after', board_to_fen(frames[i+1]['board'],frames[i+1]['side']), f'move:{i+1}')
        change(p, 'side_to_move', frames[-1]['side'], 'snapshot')
    capmap = {c['capture_id']:c for c in p.get('capture_history',[])}
    restored = 0
    for c in capmap.values():
        name = c.get('result_json')
        if name:
            path = (directory/name).resolve()
            if directory.resolve() not in path.parents:
                raise ValueError('Capture attachment is outside its game folder')
            if path.is_file():
                data = json.loads(path.read_text(encoding='utf-8'))
                observed = data.get('board')
                if observed is not None:
                    c['observed_board'] = observed
                    c['confidence_grid'] = data.get('confidence')
                    c['observation_source'] = name
                    if c.get('purpose_code') in ('initial','current_board') and observed == p.get('initial_board'):
                        c.update(accepted_board=copy.deepcopy(observed),acceptance_status='依保存初始盤面核對')
                    restored += 1
        c['expected_artifacts'] = {k:c[k] for k in ('v1_raw','v2_raw','v1_yolo','v2_yolo','result_json') if c.get(k)}
        c['artifact_status'] = {}
        for key, name in c['expected_artifacts'].items():
            path = (directory/name).resolve()
            if directory.resolve() not in path.parents:
                raise ValueError('Capture attachment is outside its game folder')
            c['artifact_status'][key] = 'available' if path.is_file() else 'not_present_in_source'
            if not path.is_file(): c[key] = None
    for i, m in enumerate(p.get('move_history', [])):
        if len(frames)>i+1 and m.get('capture_id') in capmap:
            capmap[m['capture_id']].update(accepted_board=frames[i+1]['board'],accepted_uci=m['uci'],
                                         acceptance_status='由已保存棋步重建',accepted_move_sequence=i+1)
        for motion in p.get('motion_history',[]):
            if m.get('actor')=='ai' and m.get('motion_sequence') == motion.get('sequence'):
                motion.update(verification_status='依保存棋步連結',committed_move_sequence=i+1,verification_capture_id=m.get('capture_id'))
    p['report_note'] = ('歷史報告副本；多個片段可能屬同一場，實際場數未確認。附件中的辨識盤面已補入，'
                        '無來源的耗時等數據不補猜。行棋方及 FEN 的推導修正僅存在本副本，原件未變更。'
                        f'本副本推導修正 {len(corrections)} 欄；逐欄原值見同名 report.json 的 derived_corrections。')
    p['derived_corrections'] = corrections
    return p, dict(original_issues=issues, corrections=corrections, restored_observations=restored)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--date', required=True)
    parser.add_argument('--output', required=True, type=Path)
    parser.add_argument('--correct-metadata', action='store_true')
    args = parser.parse_args()
    if not re.fullmatch(r'\d{8}',args.date): parser.error('date must be YYYYMMDD')
    if args.output.exists(): parser.error('output must be a new directory; originals are never overwritten')
    paths = sorted(CONFIG.capture_root.glob(f'{args.date}*/game.json'))
    if not paths: parser.error('no records found')
    args.output.mkdir(parents=True)
    manifest=[]; payloads=[]
    for path in paths:
        original = path.read_bytes()
        source = json.loads(original)
        p, audit = enrich(source,path.parent,correct_metadata=args.correct_metadata)
        w=GameRecordWriter._build_workbook(p)
        try: w.save(args.output/f"{p['game_id']}.xlsx")
        finally: w.close()
        (args.output/f"{p['game_id']}.report.json").write_text(json.dumps(p,ensure_ascii=False,indent=2),encoding='utf-8')
        old_book=CONFIG.record_root/f"{p['game_id']}.xlsx"
        manifest.append(dict(game_id=p['game_id'],source=str(path),json_sha256=hashlib.sha256(original).hexdigest(),
                             excel_sha256=hashlib.sha256(old_book.read_bytes()).hexdigest() if old_book.exists() else None,**audit))
        payloads.append(p)
    w=Workbook(); sheet=w.active;sheet.title='紀錄片段總覽'
    GameRecordWriter._title(sheet,'昨日紀錄片段；檔案數不代表對戰場數',8)
    sheet['A2']='使用者確認：辨識失敗後重新從目前棋盤開始，產生同場多段紀錄。實際場數未確認，未自動合併。'
    GameRecordWriter._header_row(sheet,3,['片段編號','開始時間','最後階段','提交步數','原始棋譜檢查','副本推導修正欄數','補入辨識盤面數','新版報告'])
    for i,(p,a) in enumerate(zip(payloads,manifest),4):
        values=[p['game_id'],_datetime_value(p.get('started_at')),p.get('phase'),len(p.get('move_history',[])),
                '；'.join(a['original_issues']) or '本段一致',len(a['corrections']),a['restored_observations'],'開啟片段報告']
        for col,v in enumerate(values,1): sheet.cell(i,col,v)
        sheet.cell(i,8).hyperlink=f"{p['game_id']}.xlsx";sheet.cell(i,8).style='Hyperlink'
    GameRecordWriter._finish_table(sheet,3,len(payloads)+3,[38,21,14,12,60,18,18,24])
    w.save(args.output/'紀錄片段總覽.xlsx');w.close()
    (args.output/'來源與修正清單.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    # Prove the source snapshots and workbooks stayed unchanged during export.
    for item in manifest:
        if hashlib.sha256(Path(item['source']).read_bytes()).hexdigest()!=item['json_sha256']: raise RuntimeError('source snapshot changed during report generation')
        old_book=CONFIG.record_root/f"{item['game_id']}.xlsx"
        if item['excel_sha256'] and hashlib.sha256(old_book.read_bytes()).hexdigest()!=item['excel_sha256']: raise RuntimeError('source workbook changed during report generation')
    archive=args.output.with_suffix('.zip')
    with zipfile.ZipFile(archive,'x',compression=zipfile.ZIP_DEFLATED) as z:
        for f in args.output.iterdir(): z.write(f,f.name)
    print(json.dumps(dict(fragments=len(payloads),corrected_fragments=sum(bool(m['corrections']) for m in manifest),
                         observations_restored=sum(m['restored_observations'] for m in manifest),archive=str(archive)),ensure_ascii=False))


if __name__=='__main__': main()
