"""Non-blocking, one-workbook-per-game operational records."""

from __future__ import annotations

import copy
import json
import logging
import math
import threading
import time
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .board import PIECE_LABELS
from .session import GAME_ID_PATTERN
from .record_audit import audit_history


LOGGER = logging.getLogger(__name__)

_DARK_GREEN = "1E5B4F"
_LIGHT_GREEN = "DDEBE5"
_LIGHT_GRAY = "EEF1EF"
_WHITE = "FFFFFF"
_THIN_GRAY = Side(style="thin", color="C9D2CE")


def _datetime_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return value
    # Excel does not support timezone-aware datetimes. Convert to workstation
    # local time and store a native, sortable datetime value.
    return parsed.astimezone().replace(tzinfo=None)


def _actor_label(value: Any) -> str:
    return {"player": "玩家", "ai": "AI"}.get(str(value), str(value or "—"))


def _side_label(value: Any) -> str:
    return {"w": "紅方", "b": "黑方"}.get(str(value), str(value or "—"))


def _phase_label(value: Any) -> str:
    return {
        "IDLE": "待命",
        "PREFLIGHT": "預檢",
        "INITIAL_CAPTURE": "初始拍攝",
        "PLAYER_TURN": "玩家回合",
        "AI_THINKING": "AI 計算",
        "ENGINE_READY": "等待確認引擎著法",
        "ROBOT_MOVING": "機械手臂移動",
        "VERIFYING": "盤面驗證",
        "ERROR": "錯誤",
        "FINISHED": "對局結束",
        "STOPPED": "已停止",
    }.get(str(value), str(value or "—"))


def _metric(item: dict[str, Any], key: str, *, applicable: bool = True) -> Any:
    if not applicable:
        return "不適用"
    value = item.get(key)
    return "未記錄" if value is None else value


def _detail(value: Any) -> str:
    return "未記錄" if value is None else json.dumps(value, ensure_ascii=False, separators=(",", ":"))


class GameRecordWriter:
    """Coalesce workbook updates on one background thread.

    Game and robot workflows only copy a small payload into memory. Workbook
    formatting and disk I/O never run inside the serialized hardware worker.
    """

    def __init__(self, root: Path, *, games_root: Path | None = None) -> None:
        self.root = root.resolve()
        self.games_root = games_root.resolve() if games_root is not None else None
        self._battle_ids: dict[str, str] = {}
        self._battle_heads: dict[str, dict[str, Any]] = {}
        self._condition = threading.Condition()
        self._pending: dict[str, dict[str, Any]] = {}
        self._status: dict[str, dict[str, Any]] = {}
        self._writing = False
        self._retry_after: dict[str, float] = {}
        self.retry_seconds = 5.0
        self._closing = False
        self._thread = threading.Thread(
            target=self._run,
            name="game-record-writer",
            daemon=True,
        )
        self._thread.start()

    def path_for(self, game_id: str) -> Path:
        if not GAME_ID_PATTERN.fullmatch(game_id):
            raise ValueError("game id 格式錯誤。")
        battle_id = self._battle_ids.get(game_id)
        directory = self.root / "整場對戰" if battle_id else self.root
        path = (directory / f"{battle_id or game_id}.xlsx").resolve()
        if path.parent != directory:
            raise ValueError("game id 路徑無效。")
        return path

    def schedule(self, payload: dict[str, Any]) -> None:
        game_id = str(payload.get("game_id", ""))
        self.path_for(game_id)
        snapshot = copy.deepcopy(payload)
        with self._condition:
            if self._closing:
                return
            battle_id = payload.get("battle_id")
            if self.games_root is not None and battle_id:
                if not isinstance(battle_id, str) or not GAME_ID_PATTERN.fullmatch(battle_id):
                    raise ValueError("battle id 格式錯誤。")
                self._battle_ids[game_id] = battle_id
                previous = self._battle_heads.get(battle_id, {})
                rank = lambda p: (p.get('segment_index', 1), p.get('saved_at') or '')
                if not previous or rank(snapshot) >= rank(previous):
                    self._battle_heads[battle_id] = snapshot
            # Only the newest snapshot for a game matters. This prevents Excel
            # work from building a backlog during rapid state transitions.
            self._pending[game_id] = snapshot
            self._retry_after.pop(game_id, None)
            self._status.setdefault(game_id, {}).update(pending=True)
            self._condition.notify_all()

    def status_for(self, game_id: str) -> dict[str, Any]:
        self.path_for(game_id)
        with self._condition:
            return copy.deepcopy(self._status.get(game_id, {}))

    def flush(self, timeout: float = 5.0) -> bool:
        deadline = time.monotonic() + timeout
        with self._condition:
            while self._pending or self._writing:
                remaining = deadline - time.monotonic()
                if remaining <= 0:
                    return False
                self._condition.wait(remaining)
            return True

    def close(self, timeout: float = 5.0) -> None:
        with self._condition:
            self._closing = True
            self._condition.notify_all()
        self._thread.join(timeout=timeout)

    def _run(self) -> None:
        while True:
            with self._condition:
                self._condition.wait_for(lambda: self._pending or self._closing)
                if not self._pending and self._closing:
                    self._condition.notify_all()
                    return
                now = time.monotonic()
                ready = [key for key in self._pending if self._closing or self._retry_after.get(key, 0) <= now]
                if not ready:
                    self._condition.wait(timeout=max(0.01, min(self._retry_after[key] for key in self._pending) - now))
                    continue
                game_id = ready[0]
                payload = self._pending[game_id]
                del self._pending[game_id]
                self._writing = True
            try:
                self._write(payload)
            except Exception as exc:
                # Recording must never unlock, replay, or delay a robot action.
                # The durable game.json remains the recovery source of truth.
                LOGGER.exception("[RECORD] Excel update failed game=%s", game_id)
                with self._condition:
                    if isinstance(exc, OSError) and not self._closing:
                        self._pending.setdefault(game_id, payload)
                        self._retry_after[game_id] = time.monotonic() + self.retry_seconds
                    self._status.setdefault(game_id, {}).update(
                        pending=game_id in self._pending,
                        last_error=str(exc),
                        last_error_at=datetime.now().astimezone().isoformat(),
                    )
            else:
                with self._condition:
                    self._status.setdefault(game_id, {}).update(
                        pending=game_id in self._pending,
                        last_error=None,
                        last_success_at=datetime.now().astimezone().isoformat(),
                    )
            finally:
                with self._condition:
                    self._writing = False
                    self._condition.notify_all()

    def _write(self, payload: dict[str, Any]) -> Path:
        # Excel is optional: an unavailable export directory must not prevent
        # the UI and authoritative game session service from starting.
        self.root.mkdir(parents=True, exist_ok=True)
        game_id = str(payload["game_id"])
        target = self.path_for(game_id)
        temporary = target.with_name(f".{game_id}.tmp.xlsx")
        target.parent.mkdir(parents=True, exist_ok=True)
        battle_id = self._battle_ids.get(game_id)
        if battle_id:
            with self._condition:
                head = copy.deepcopy(self._battle_heads[battle_id])
            workbook = self._build_battle_workbook(self._battle_segments(head))
        else:
            workbook = self._build_workbook(payload)
        try:
            workbook.save(temporary)
            for attempt in range(4):
                try:
                    temporary.replace(target)
                    break
                except PermissionError:
                    if attempt == 3:
                        raise
                    # Windows virus scanning and file indexing can briefly
                    # hold a newly written XLSX. Retry only in this background
                    # thread; never delay or replay a robot command.
                    time.sleep(0.05 * (2 ** attempt))
        finally:
            workbook.close()
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                LOGGER.warning("[RECORD] could not remove temporary workbook %s", temporary, exc_info=True)
        LOGGER.info("[RECORD] Excel updated game=%s path=%s", game_id, target)
        return target

    def _battle_segments(self, payload: dict[str, Any]) -> list[dict[str, Any]]:
        """Read only explicitly linked predecessors, never infer a battle by time."""
        chain = [payload]
        seen = {payload["game_id"]}
        root_id = payload.get("battle_id") or payload["game_id"]
        while chain[-1].get("parent_game_id"):
            parent = chain[-1]["parent_game_id"]
            if not isinstance(parent, str) or not GAME_ID_PATTERN.fullmatch(parent) or parent in seen:
                raise ValueError("對戰片段關聯無效或形成循環。")
            if self.games_root is None or len(chain) >= 1000:
                raise ValueError("無法讀取對戰前段紀錄。")
            path = (self.games_root / parent / "game.json").resolve()
            if self.games_root not in path.parents:
                raise ValueError("對戰片段路徑無效。")
            previous = json.loads(path.read_text(encoding="utf-8"))
            if previous.get("game_id") != parent or (previous.get("battle_id") or parent) != root_id:
                raise ValueError("對戰前段編號與來源不一致。")
            seen.add(parent)
            chain.append(previous)
        return list(reversed(chain))

    @classmethod
    def _build_battle_workbook(cls, segments: list[dict[str, Any]]) -> Workbook:
        """One workbook for a battle; keep each baseline and history auditable."""
        workbook = cls._build_workbook(segments[-1])
        summary = workbook["對局摘要"]
        for row in summary:
            counts = {'已提交著法': 'move_history', '拍攝次數': 'capture_history', 'Robot 動作次數': 'motion_history', '引擎查詢次數': 'engine_history'}
            if row[0].value in counts:
                row[1].value = sum(len(p.get(counts[row[0].value], [])) for p in segments)
            if row[0].value == "歷史一致性":
                row[1].value = '；'.join(f"片段 {i}：{'／'.join(audit_history(p)[1]) or '本段一致'}" for i,p in enumerate(segments,1))
            if row[0].value == "棋譜銜接" and any(p.get("history_gap") for p in segments):
                row[1].value = "含重新建立基準的片段；各段分別重播，中間未知棋步未補猜。"
        for name, key, write in (("著法紀錄", "move_history", cls._write_moves), ("拍攝紀錄", "capture_history", cls._write_captures),
                                 ("Robot動作", "motion_history", cls._write_motions), ("流程紀錄", "event_history", cls._write_events),
                                 ("引擎紀錄", "engine_history", cls._write_engine)):
            old = workbook[name]; position = workbook.index(old); workbook.remove(old)
            sheet = workbook.create_sheet(name, position)
            items = [item for p in segments for item in p.get(key, [])]
            write(sheet, items)
            column = sheet.max_column + 1
            cls._header_row(sheet, 3, [cell.value for cell in sheet[3]] + ["來源片段", "片段內序號"])
            row = 4
            for p in segments:
                for item in p.get(key, []):
                    sheet.cell(row, column, p['game_id']); sheet.cell(row, column + 1, item.get('sequence'))
                    row += 1
            sheet.column_dimensions[get_column_letter(column)].width = 38
            sheet.column_dimensions[get_column_letter(column+1)].width = 16
            sheet.auto_filter.ref = f"A3:{get_column_letter(column+1)}{max(3,row-1)}"
        old = workbook['棋盤履歷']; workbook.remove(old)
        old_observations = workbook['辨識盤面']; workbook.remove(old_observations)
        observations = workbook.create_sheet('辨識盤面')
        cls._write_observations(observations,[dict(c,_source_segment=p['game_id']) for p in segments for c in p.get('capture_history',[])])
        for index, p in enumerate(segments, 1):
            sheet = workbook.create_sheet('棋盤履歷' if len(segments)==1 else f'棋盤履歷{index}')
            cls._write_boards(sheet, p)
            sheet['A2'] = f"來源片段：{p['game_id']}；本段獨立重播，不假定跨段棋步完整。"
        sheet = workbook.create_sheet('對戰片段', 1)
        cls._title(sheet, '同一場對戰的紀錄片段', 8)
        cls._header_row(sheet, 3, ['片段', '紀錄編號', '開始時間', '最後狀態', '已提交步數', '重新建立基準', '歷史一致性', '最後訊息'])
        for i, p in enumerate(segments, 1):
            _, issues = audit_history(p)
            values = [i,p['game_id'],_datetime_value(p.get('started_at')),_phase_label(p.get('phase')),len(p.get('move_history',[])),
                      '有棋譜缺口' if p.get('history_gap') else '—','；'.join(issues) if issues else '本段一致',p.get('message')]
            for col,value in enumerate(values,1): sheet.cell(i+3,col,value)
        cls._finish_table(sheet,3,len(segments)+3,[8,38,21,14,14,20,45,55])
        for sheet in workbook:
            sheet.freeze_panes = sheet.freeze_panes or 'A4'
            sheet.print_area = sheet.dimensions
            sheet.page_setup.orientation = 'landscape'
            for row in sheet:
                for cell in row:
                    if cell.data_type == 'f': cell.data_type = 's'
        return workbook

    @classmethod
    def _build_workbook(cls, payload: dict[str, Any]) -> Workbook:
        workbook = Workbook()
        summary = workbook.active
        summary.title = "對局摘要"
        moves = workbook.create_sheet("著法紀錄")
        captures = workbook.create_sheet("拍攝紀錄")
        motions = workbook.create_sheet("Robot動作")
        events = workbook.create_sheet("流程紀錄")
        boards = workbook.create_sheet("棋盤履歷")
        engine = workbook.create_sheet("引擎紀錄")
        observations = workbook.create_sheet("辨識盤面")

        cls._write_summary(summary, payload)
        cls._write_moves(moves, payload.get("move_history", []))
        cls._write_captures(captures, payload.get("capture_history", []))
        cls._write_motions(motions, payload.get("motion_history", []))
        cls._write_events(events, payload.get("event_history", []))
        cls._write_boards(boards, payload)
        cls._write_engine(engine, payload.get("engine_history", []))
        cls._write_observations(observations, payload.get("capture_history", []))
        for sheet in workbook:
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.print_options.horizontalCentered = True
            sheet.print_area = sheet.dimensions
            sheet.page_setup.orientation = "landscape"
            # Wide diagnostic tables span pages instead of unreadable shrinking.
            sheet.page_setup.fitToWidth = 0
            sheet.page_setup.fitToHeight = 0
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f":
                        # Logs and model messages are literal evidence, never formulas.
                        cell.data_type = "s"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
        return workbook

    @classmethod
    def _title(cls, sheet, text: str, end_column: int) -> None:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
        cell = sheet.cell(1, 1, text)
        cell.fill = PatternFill("solid", fgColor=_DARK_GREEN)
        cell.font = Font(color=_WHITE, bold=True, size=16)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[1].height = 28

    @classmethod
    def _section(cls, sheet, row: int, text: str, end_column: int) -> None:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_column)
        cell = sheet.cell(row, 1, text)
        cell.fill = PatternFill("solid", fgColor=_LIGHT_GREEN)
        cell.font = Font(bold=True, color=_DARK_GREEN)

    @classmethod
    def _write_summary(cls, sheet, payload: dict[str, Any]) -> None:
        cls._title(sheet, "TM5-700 象棋機器人｜逐局紀錄", 9)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A4"
        sheet["A2"] = "時間欄位為工作站本地時間；本檔由系統背景更新，不參與機械手臂動作判斷。"
        sheet.merge_cells("A2:I2")
        sheet["A2"].font = Font(italic=True, color="53645D")

        state_rows = [
            ("對局編號", payload.get("game_id")),
            ("對戰編號", payload.get("battle_id") or "未建立關聯（舊紀錄不能直接計為一場）"),
            ("上一個片段", payload.get("parent_game_id") or "—"),
            ("同場片段序號", payload.get("segment_index") or "未記錄"),
            ("建立片段原因", {"new_battle": "全新對戰", "rebaseline_after_interruption": "中斷後重新建立基準"}.get(payload.get("segment_reason"), "舊紀錄未記錄")),
            ("棋譜銜接", "重新建立基準；中間著法未確認，不能視為連續棋譜" if payload.get("history_gap") else "本片段內重播；跨片段仍須核對"),
            ("紀錄類型", "尚未建立盤面" if payload.get("initial_board") is None else
             "載入終局盤面（不另計完成對戰）" if payload.get("phase") == "FINISHED" and not payload.get("move_history") else "對戰片段（不能以檔案數計場次）"),
            ("開始時間", _datetime_value(payload.get("started_at"))),
            ("最後更新", _datetime_value(payload.get("saved_at"))),
            ("對局狀態", _phase_label(payload.get("phase"))),
            ("目前回合", _actor_label(payload.get("current_turn"))),
            ("先手", _actor_label(payload.get("first_actor"))),
            ("玩家棋色", _side_label(payload.get("player_side"))),
            ("AI 棋色", _side_label(payload.get("ai_side"))),
            ("對局模式", {"current_board": "從目前棋盤", "human_red": "玩家執紅", "human_black": "玩家執黑", "engine_vs_engine": "AI 對 AI（逐步確認）", "analysis": "雙人分析"}.get(payload.get("game_mode"), "標準開局")),
            ("應行棋方", _side_label(payload.get("side_to_move"))),
            ("紅方控制者", {"human": "玩家", "engine": "Pikafish"}.get(payload.get("controllers", {}).get("w"), "未記錄")),
            ("黑方控制者", {"human": "玩家", "engine": "Pikafish"}.get(payload.get("controllers", {}).get("b"), "未記錄")),
            ("已提交著法", len(payload.get("move_history", []))),
            ("拍攝次數", len(payload.get("capture_history", []))),
            ("Robot 動作次數", len(payload.get("motion_history", []))),
            ("最近辨識編號", payload.get("last_capture_id") or "—"),
            ("待確認 AI 著法", (payload.get("pending_move") or {}).get("uci") or "—"),
            ("重試次數", payload.get("pending_retry_count", 0)),
            ("動作狀態未知", "是" if payload.get("motion_state_uncertain") else "否"),
            ("狀態訊息", payload.get("message") or "—"),
            ("引擎查詢次數", len(payload.get("engine_history", [])) if "engine_history" in payload else "未記錄"),
            ("空白數據說明", "未記錄＝來源沒有數據；不適用＝該角色無此步驟。不可當作零。流程總秒數含等待／恢復時間。"),
            ("報告副本說明", payload.get("report_note") or "執行中紀錄；以保存快照為來源。"),
        ]
        cls._section(sheet, 4, "對局狀態", 9)
        row = 5
        for label, value in state_rows:
            sheet.cell(row, 1, label)
            sheet.cell(row, 2, value)
            sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=9)
            if isinstance(value, str):
                units = sum(2 if unicodedata.east_asian_width(c) in ('W','F') else 1 for c in value)
                sheet.row_dimensions[row].height = min(409, max(20, ((units + 94) // 95) * 16 + 4))
            row += 1

        _, issues = audit_history(payload)
        sheet.cell(row, 1, "歷史一致性")
        sheet.cell(row, 2, "；".join(issues) if issues else "盤面、行棋方及已記錄著法一致（不代表真機驗證）")
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=9)
        sheet.row_dimensions[row].height = max(32, 16 * (len(issues) + 1))
        row += 1

        settings = payload.get("settings", {})
        hardware_enabled = settings.get("hardware_enabled")
        cls._section(sheet, row, "引擎與真機設定快照", 9)
        row += 1
        setting_rows = [
            ("執行模式", "真機" if hardware_enabled is True else "未啟用" if hardware_enabled is False else "未記錄"),
            ("Pikafish 深度", settings.get("engine_depth")),
            ("Pikafish 思考時間（ms）", settings.get("engine_movetime_ms")),
            ("Robot Base / TCP", f"{settings.get('base_name', '—')} / {settings.get('tcp_name', '—')}"),
            ("移動速度 / 加速度", f"{settings.get('speed', '—')} / {settings.get('acceleration', '—')}"),
            ("Z_SAFE / Z_APPROACH / Z_PICK", f"{settings.get('z_safe', '—')} / {settings.get('z_approach', '—')} / {settings.get('z_pick', '—')}"),
            ("影像信心門檻", settings.get("vision_confidence")),
        ]
        for label, value in setting_rows:
            sheet.cell(row, 1, label)
            sheet.cell(row, 2, value if value is not None else "—")
            sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=9)
            row += 1

        row += 1
        cls._section(sheet, row, "最後保存盤面（程式座標：第 0 列為紅方底線）", 9)
        row += 1
        for column in range(9):
            cell = sheet.cell(row, column + 1, chr(ord("A") + column))
            cell.fill = PatternFill("solid", fgColor=_DARK_GREEN)
            cell.font = Font(color=_WHITE, bold=True)
            cell.alignment = Alignment(horizontal="center")
        board = payload.get("board")
        if isinstance(board, list):
            for board_row, values in enumerate(board):
                for column, piece in enumerate(values):
                    label = ""
                    if isinstance(piece, int) and piece:
                        label = f"{PIECE_LABELS.get(piece, '?')} ({piece})"
                    cell = sheet.cell(row + board_row + 1, column + 1, label)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = Border(left=_THIN_GRAY, right=_THIN_GRAY, top=_THIN_GRAY, bottom=_THIN_GRAY)

        sheet.column_dimensions["A"].width = 24
        for column in range(2, 10):
            sheet.column_dimensions[get_column_letter(column)].width = 14
        for cells in sheet.iter_rows(min_row=5, max_row=row - 2, min_col=1, max_col=1):
            cells[0].font = Font(bold=True, color="33443D")
            cells[0].fill = PatternFill("solid", fgColor=_LIGHT_GRAY)
        for cells in sheet.iter_rows():
            for cell in cells:
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                cell.alignment = Alignment(
                    horizontal=cell.alignment.horizontal or "left",
                    vertical="center",
                    wrap_text=True,
                )

    @classmethod
    def _write_moves(cls, sheet, history: Any) -> None:
        headers = [
            "序號", "輪次", "提交時間", "行動者", "棋色", "UCI",
            "起點棋盤座標", "起點矩陣", "起點玩家矩陣", "起點P編號", "起點Robot X", "起點Robot Y",
            "終點棋盤座標", "終點矩陣", "終點玩家矩陣", "終點P編號", "終點Robot X", "終點Robot Y",
            "是否吃子", "被吃棋子", "FEN 前", "FEN 後", "FEN錯誤", "確認影像編號", "影像修正格",
            "拍攝秒數", "辨識秒數", "影像總秒數", "合法性檢查秒數", "引擎秒數", "Robot秒數", "驗證秒數", "本回合秒數", "Robot動作序號",
            "著法作業編號", "流程開始時間",
        ]
        cls._title(sheet, "著法、座標、FEN 與耗時", len(headers))
        cls._header_row(sheet, 3, headers)
        rows = history if isinstance(history, list) else []
        for index, item in enumerate(rows, start=1):
            captured_id = item.get("captured_piece_id", 0)
            source = item.get("source_position") if isinstance(item.get("source_position"), dict) else {}
            destination = item.get("destination_position") if isinstance(item.get("destination_position"), dict) else {}
            values = [
                index,
                item.get("round"),
                _datetime_value(item.get("committed_at")),
                _actor_label(item.get("actor")),
                _side_label(item.get("side")),
                item.get("uci") or "—",
                source.get("board_text") or item.get("source_text") or "—",
                source.get("matrix") or "—",
                source.get("player_matrix") or "—",
                source.get("player_point"),
                source.get("robot_x"),
                source.get("robot_y"),
                destination.get("board_text") or item.get("destination_text") or "—",
                destination.get("matrix") or "—",
                destination.get("player_matrix") or "—",
                destination.get("player_point"),
                destination.get("robot_x"),
                destination.get("robot_y"),
                "是" if item.get("captured") else "否",
                PIECE_LABELS.get(captured_id, "—") if captured_id else "—",
                item.get("fen_before") or "—",
                item.get("fen_after") or "—",
                item.get("fen_error") or "—",
                item.get("capture_id") or "—",
                ", ".join(item.get("vision_repaired_cells", [])) or "—",
                _metric(item, "capture_seconds"),
                _metric(item, "vision_processing_seconds"),
                _metric(item, "vision_total_seconds"),
                _metric(item, "legality_seconds"),
                _metric(item, "engine_seconds", applicable=item.get("actor") == "ai"),
                _metric(item, "robot_seconds", applicable=item.get("actor") == "ai"),
                _metric(item, "verification_seconds", applicable=item.get("actor") == "ai"),
                _metric(item, "turn_seconds"),
                _metric(item, "motion_sequence", applicable=item.get("actor") == "ai"),
                item.get("move_attempt_id") or "未記錄", _datetime_value(item.get("turn_started_at")) or "未記錄",
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(index + 3, column, value)
        cls._finish_table(sheet, 3, len(rows) + 3, [8, 8, 21, 10, 10, 12] + [14] * 12 + [10, 12, 34, 34, 28, 31, 20] + [13] * 9 + [36, 21])

    @classmethod
    def _write_captures(cls, sheet, history: Any) -> None:
        headers = [
            "序號", "輪次", "開始時間", "完成時間", "用途", "狀態", "影像編號",
            "V1原圖", "V2原圖", "V1辨識圖", "V2辨識圖", "結果JSON",
            "收圖秒數", "辨識秒數", "總秒數", "棋子數", "最低信心", "平均信心", "最高信心",
            "錯誤類型", "錯誤訊息",
            "原始盤面檢查", "沿用歷史格數", "歷史修正明細",
            "原始辨識盤面", "採用盤面", "逐格信心", "採用狀態", "採用著法", "失敗階段", "檔案狀態", "流程經過秒數",
        ]
        cls._title(sheet, "拍攝與辨識紀錄", len(headers))
        cls._header_row(sheet, 3, headers)
        rows = history if isinstance(history, list) else []
        for index, item in enumerate(rows, start=1):
            values = [
                index, item.get("round"), _datetime_value(item.get("started_at")),
                _datetime_value(item.get("completed_at")), item.get("purpose"), item.get("status"),
                item.get("capture_id"), item.get("v1_raw") or "未收到／未記錄", item.get("v2_raw") or "未收到／未記錄",
                item.get("v1_yolo") or "未產生／未記錄", item.get("v2_yolo") or "未產生／未記錄", item.get("result_json") or "未產生／未記錄",
                _metric(item, "capture_seconds"), _metric(item, "processing_seconds"), _metric(item, "total_seconds"),
                _metric(item, "detected_pieces"), _metric(item, "confidence_min"), _metric(item, "confidence_avg"),
                _metric(item, "confidence_max"), item.get("error_type") or "—", item.get("error_message") or "—",
                item.get("validation_error") or "—", len(item.get("history_repairs", [])),
                "、".join(f"{repair['cell']}：{repair.get('observed')}→{repair.get('restored')}" for repair in item.get("history_repairs", [])) or "—",
                f"見辨識盤面第 {index} 筆" if item.get('observed_board') is not None else '未記錄',
                f"見辨識盤面第 {index} 筆" if item.get('accepted_board') is not None else '未採用／未記錄',
                f"見辨識盤面第 {index} 筆" if item.get('confidence_grid') is not None else '未記錄',
                item.get("acceptance_status") or "未提交著法", item.get("accepted_uci") or "—",
                item.get("failure_stage") or "—", _detail(item.get("artifact_status")),
                _metric(item,'workflow_seconds'),
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(index + 3, column, value)
        cls._finish_table(sheet, 3, len(rows) + 3, [8, 8, 21, 21, 20, 12, 31] + [34] * 5 + [12] * 7 + [20, 42, 40, 12, 48] + [30, 30, 30, 18, 12, 26, 48, 16])

    @classmethod
    def _write_motions(cls, sheet, history: Any) -> None:
        headers = [
            "序號", "輪次", "開始時間", "完成時間", "狀態", "重試", "UCI", "耗時秒數",
            "起點矩陣", "起點P編號", "起點X", "起點Y", "終點矩陣", "終點P編號", "終點X", "終點Y",
            "目標X", "目標Y", "目標Z", "實際X", "實際Y", "實際Z", "XYZ誤差mm", "姿態誤差deg",
            "落子箱X", "落子箱Y", "錯誤類型", "錯誤訊息",
            "著法作業編號", "落子影像驗證", "提交棋步序號", "確認影像", "最後動作階段", "動作階段明細",
        ]
        cls._title(sheet, "Robot 動作與位置誤差", len(headers))
        cls._header_row(sheet, 3, headers)
        rows = history if isinstance(history, list) else []
        for index, item in enumerate(rows, start=1):
            source = item.get("source") if isinstance(item.get("source"), dict) else {}
            destination = item.get("destination") if isinstance(item.get("destination"), dict) else {}
            values = [
                index, item.get("round"), _datetime_value(item.get("started_at")),
                _datetime_value(item.get("completed_at")), item.get("status"), "是" if item.get("retry") else "否",
                item.get("move_uci"), item.get("duration_seconds"), source.get("matrix"), source.get("player_point"),
                source.get("robot_x"), source.get("robot_y"), destination.get("matrix"), destination.get("player_point"),
                destination.get("robot_x"), destination.get("robot_y"), item.get("target_x"), item.get("target_y"),
                item.get("target_z"), item.get("actual_x"), item.get("actual_y"), item.get("actual_z"),
                item.get("position_error_mm"), item.get("orientation_error_deg"), item.get("graveyard_x"),
                item.get("graveyard_y"), item.get("error_type") or "—", item.get("error_message") or "—",
                item.get("move_attempt_id") or "未記錄", item.get("verification_status") or "未記錄",
                item.get("committed_move_sequence") or "尚未連結提交", item.get("verification_capture_id") or "—",
                item.get("last_action") or "未記錄", _detail(item.get("stages")),
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(index + 3, column, value)
        cls._finish_table(sheet, 3, len(rows) + 3, [8, 8, 21, 21, 12, 8, 12, 12] + [13] * 18 + [20, 42] + [36, 20, 16, 31, 30, 48])

    @classmethod
    def _write_engine(cls, sheet, history: Any) -> None:
        headers = ["序號", "開始時間", "完成時間", "用途", "狀態", "棋色", "耗時秒數", "待驗著法", "最佳著法",
                   "結果", "對局深度設定", "思考時間設定ms", "歷史步數", "FEN", "候選路線與分數", "錯誤類型", "錯誤訊息",
                   "主路線深度", "分數類型", "分數", "紅勝估計‰", "和棋估計‰", "黑勝估計‰"]
        cls._title(sheet, "引擎查詢、局勢分析與失敗紀錄", len(headers))
        cls._header_row(sheet, 3, headers)
        rows = history if isinstance(history, list) else []
        for index, item in enumerate(rows, 1):
            lines = item.get('lines') or []
            primary = next((line for line in lines if line.get('multipv') == 1), {})
            description = '\n'.join(f"路線 {line.get('multipv')} 深度 {line.get('depth')} 分數 {line.get('score_type')}:{line.get('score')}；"
                                    f"nodes={line.get('nodes')} nps={line.get('nps')} time_ms={line.get('time_ms')}；"
                                    + ' '.join(line.get('pv',[])) for line in lines) or '未記錄'
            values = [index, _datetime_value(item.get("started_at")), _datetime_value(item.get("completed_at")),
                      item.get("operation"), item.get("status"), _side_label(item.get("side")), _metric(item, "duration_seconds"),
                      item.get("requested_move") or "—", item.get("bestmove") or "—", _detail(item.get("result")),
                      item.get("depth"), item.get("movetime_ms"), item.get("history_length"), item.get("fen"),
                      description, item.get("error_type") or "—", item.get("error_message") or "—",
                      *[_metric(primary,k) for k in ('depth','score_type','score','red_win_per_mille','draw_per_mille','black_win_per_mille')]]
            for column, value in enumerate(values, 1):
                sheet.cell(index + 3, column, value)
        cls._finish_table(sheet, 3, len(rows) + 3, [8,21,21,18,14,10,14,12,12,18,12,18,12,45,60,22,48]+[16]*6)

    @classmethod
    def _write_observations(cls, sheet, history: Any) -> None:
        cls._title(sheet, '拍攝時的原始辨識、採用盤面與逐格信心', 30)
        sheet['A2'] = '第 9 列在上、第 0 列在下。未取得代表沒有來源；· 代表空格，空格分類信心不適用。信心值不是正確率。'
        sheet.merge_cells('A2:AD2')
        sheet.row_dimensions[2].height = 28
        sheet.freeze_panes = 'B4'; sheet.sheet_view.showGridLines = False
        rows = history if isinstance(history,list) else []
        if not rows: sheet['A4']='沒有拍攝紀錄。'
        for index,item in enumerate(rows,1):
            top=4+(index-1)*14
            sheet.cell(top,1,f"第 {index} 筆　{item.get('capture_id','—')}　{item.get('status','—')}　{item.get('_source_segment','')}")
            sheet.merge_cells(start_row=top,start_column=1,end_row=top,end_column=30)
            sheet.cell(top,1).fill=PatternFill('solid',fgColor=_LIGHT_GREEN)
            for start,key,label in [(2,'observed_board','原始辨識'),(12,'accepted_board','採用盤面'),(22,'confidence_grid','逐格信心')]:
                sheet.cell(top+1,start,label)
                sheet.merge_cells(start_row=top+1,start_column=start,end_row=top+1,end_column=start+8)
                grid=item.get(key)
                for x in range(9): sheet.cell(top+2,start+x,chr(97+x))
                for offset,y in enumerate(range(9,-1,-1)):
                    row=top+3+offset
                    sheet.cell(row,1,y);sheet.row_dimensions[row].height=22
                    for x in range(9):
                        value=grid[y][x] if isinstance(grid,list) and len(grid)==10 and len(grid[y])==9 else None
                        text='未取得' if value is None else value if key=='confidence_grid' else PIECE_LABELS.get(value,'·')
                        observed=item.get('observed_board')
                        if key=='confidence_grid' and grid is not None and value is None and observed is not None and observed[y][x]==0:
                            text='不適用'
                        cell=sheet.cell(row,start+x,text)
                        cell.alignment=Alignment(horizontal='center',vertical='center')
                        cell.border=Border(left=_THIN_GRAY,right=_THIN_GRAY,top=_THIN_GRAY,bottom=_THIN_GRAY)
                        if key=='confidence_grid' and isinstance(value,(int,float)): cell.number_format='0.0%'
                        elif isinstance(value,int) and value>=8: cell.font=Font(color='B42318')
        sheet.column_dimensions['A'].width=6
        for col in range(2,31): sheet.column_dimensions[get_column_letter(col)].width=10
        for col in ('K','U'): sheet.column_dimensions[col].width=3

    @classmethod
    def _write_events(cls, sheet, history: Any) -> None:
        cls._title(sheet, "流程與異常紀錄", 14)
        headers = ["序號", "時間", "階段", "目前回合", "訊息", "影像編號", "目前動作", "下一動作", "動作狀態未知", "錯誤代碼", "錯誤類型", "錯誤訊息", "事件種類", "確認者"]
        cls._header_row(sheet, 3, headers)
        rows = history if isinstance(history, list) else []
        for index, item in enumerate(rows, start=1):
            error = item.get("error") if isinstance(item.get("error"), dict) else {}
            values = [
                index,
                _datetime_value(item.get("recorded_at")),
                _phase_label(item.get("phase")),
                _actor_label(item.get("current_turn")),
                item.get("message") or "—",
                item.get("capture_id") or "—",
                item.get("current_action") or "—",
                item.get("next_action") or "—",
                "是" if item.get("motion_state_uncertain") else "否",
                error.get("code") or "—",
                error.get("type") or "—",
                error.get("message") or "—", item.get("event_type") or "狀態變更", item.get("confirmed_by") or "—",
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(index + 3, column, value)
        cls._finish_table(sheet, 3, len(rows) + 3, [8, 21, 14, 11, 42, 31, 28, 28, 14, 22, 20, 48, 28, 22])

    @staticmethod
    def _header_row(sheet, row: int, headers: list[str]) -> None:
        for column, value in enumerate(headers, start=1):
            cell = sheet.cell(row, column, value)
            cell.fill = PatternFill("solid", fgColor=_DARK_GREEN)
            cell.font = Font(color=_WHITE, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    @staticmethod
    def _finish_table(sheet, header_row: int, final_row: int, widths: list[int]) -> None:
        sheet.freeze_panes = f"C{header_row + 1}"
        sheet.print_title_rows = f"1:{header_row}"
        sheet.row_dimensions[header_row].height = 32
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(widths))}{max(header_row, final_row)}"
        sheet.sheet_view.showGridLines = False
        for column, width in enumerate(widths, start=1):
            # A timestamp must not run into the adjacent actor/status column.
            if width == 21:
                width = 24
            sheet.column_dimensions[get_column_letter(column)].width = width
        for row in sheet.iter_rows(min_row=header_row + 1, max_row=final_row):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                elif isinstance(cell.value, float):
                    cell.number_format = "0.000"
        for row in range(header_row + 1, final_row + 1):
            lines = 1
            for cell in sheet[row]:
                if cell.value is None or isinstance(cell.value, (int, float, datetime)):
                    continue
                width = sheet.column_dimensions[cell.column_letter].width
                wrapped = sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(ch) in 'WF' else 1 for ch in line) / max(width - 2, 1))) for line in str(cell.value).splitlines())
                lines = max(lines, wrapped)
            sheet.row_dimensions[row].height = min(409, max(30, lines * 15 + 6))
            if row % 2 == 1:
                for cell in sheet[row]:
                    cell.fill = PatternFill("solid", fgColor="F6F8F7")

    @classmethod
    def _write_boards(cls, sheet, payload):
        cls._title(sheet, "逐步棋盤（紅方在下、黑方在上）", 10)
        sheet.merge_cells("A2:J2")
        sheet["A2"] = "由初始盤面與已提交著法重播。黃色為該步起終點；待確認 AI 著法不列入。"
        sheet.row_dimensions[2].height = 30
        sheet["A2"].alignment = Alignment(wrap_text=True)
        frames, issues = audit_history(payload)
        sheet.merge_cells("A3:J3")
        sheet["A3"] = "；".join(issues) if issues else "重播與保存盤面、行棋方一致。"
        sheet["A3"].alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[3].height = max(30, 18 * len(issues))
        sheet["A3"].font = Font(color="B42318" if issues else "33443D", size=11)
        sheet.column_dimensions["A"].width = 7
        for c in range(2, 11):
            sheet.column_dimensions[get_column_letter(c)].width = 9
        sheet.freeze_panes = "B4"
        sheet.sheet_view.showGridLines = False
        for index, frame in enumerate(frames):
            top = 5 + index * 14
            cls._section(sheet, top, f"第 {frame['sequence']} 步　{frame['uci']}　下一方：{_side_label(frame['side'])}", 10)
            for c in range(9):
                sheet.cell(top + 1, c + 2, chr(97 + c))
            for offset, rank in enumerate(range(9, -1, -1)):
                row = top + 2 + offset
                sheet.cell(row, 1, rank)
                sheet.row_dimensions[row].height = 25
                for c, piece in enumerate(frame["board"][rank]):
                    cell = sheet.cell(row, c + 2, PIECE_LABELS.get(piece, "·"))
                    cell.font = Font(name="Microsoft JhengHei", size=12, color="B42318" if piece >= 8 else "202020")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = Border(left=_THIN_GRAY, right=_THIN_GRAY, top=_THIN_GRAY, bottom=_THIN_GRAY)
                    if (c, rank) in (frame.get("source"), frame.get("destination")):
                        cell.fill = PatternFill("solid", fgColor="FFF0B3")
