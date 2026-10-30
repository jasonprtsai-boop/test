import json
import tempfile
import threading
import unittest
from unittest.mock import patch
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook, Workbook

from tmchess.board import Move, apply_move
from tmchess.config import CONFIG, SITE_ROOT
from tmchess.records import GameRecordWriter
from tmchess.session import GameSessionStore
from tmchess.state import Phase
from tmchess.web import create_ui_app
from tests.fakes import make_service, starting_board, wait_for


class GameRecordTests(unittest.TestCase):
    def test_photo_archive_matches_workbook_and_combines_segments(self):
        self.records.games_root = self.sessions.root
        ids = ['20261028T010000Z-12345678', '20261028T020000Z-87654321']
        for index, game_id in enumerate(ids):
            directory = self.sessions.create(game_id)
            (directory / f'capture_{index}_v1.jpg').write_bytes(bytes([index, 1, 2]))
        self.records._battle_ids[ids[1]] = ids[0]
        self.records._record_names[ids[0]] = '測試者_'
        target = self.records.path_for(ids[1])
        segments = [{'game_id': game_id} for game_id in ids]
        self.records._export_photos(target, segments)
        self.assertEqual(target.stem, '測試者_2026-10-28_09-00-00_12345678')
        destination = self.records.root.parent / '遊戲紀錄照片' / target.stem
        self.assertEqual(len(list(destination.glob('*.jpg'))), 2)
        self.assertFalse(target.with_suffix('').exists())
        photo = destination / 'capture_0_v1.jpg'
        before = photo.stat().st_mtime_ns
        self.records._export_photos(target, segments)
        self.assertEqual(photo.stat().st_mtime_ns, before)
        self.assertEqual(photo.read_bytes(), bytes([0, 1, 2]))
        self.assertTrue((self.sessions.game_directory(ids[0]) / photo.name).is_file())

    def test_record_date_folder_uses_taiwan_day(self):
        path = self.records.path_for('20261022T163000Z-12345678')
        self.assertEqual(path.parent, self.records.root / '2026-10-23')

    def test_diagnostic_columns_keep_missing_values_explicit(self):
        workbook = Workbook()
        captures = workbook.active
        GameRecordWriter._write_captures(captures, [{
            'trigger_seconds': 5.2, 'stationary_confirmation_seconds': 1.3,
            'repair_context': '玩家走子', 'failure_stage': 'capture_trigger',
        }, {}])
        headers = {cell.value: cell.column for cell in captures[3]}
        self.assertEqual(captures.cell(4, headers['拍攝觸發秒數']).value, 5.2)
        self.assertEqual(captures.cell(4, headers['靜止確認秒數']).value, 1.3)
        self.assertEqual(captures.cell(5, headers['拍攝觸發秒數']).value, '未記錄')
        motions = workbook.create_sheet()
        GameRecordWriter._write_motions(motions, [{'confirmation_method': 'motion_completed'}])
        headers = {cell.value: cell.column for cell in motions[3]}
        self.assertIn('未做影像驗證', motions.cell(4, headers['棋步確認依據']).value)
        self.assertIn('不能推定', motions.cell(4, headers['人工補放紀錄']).value)
        workbook.close()

    def setUp(self):
        (SITE_ROOT / "runtime").mkdir(exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(dir=SITE_ROOT / "runtime")
        root = Path(self.temporary.name)
        self.sessions = GameSessionStore(root / "games")
        self.records = GameRecordWriter(root / "對局紀錄")
        initial = starting_board()
        player_move = Move((0, 3), (0, 4), 0, "a3a4")
        after_player = apply_move(initial, player_move)
        ai_move = Move((0, 6), (0, 5), 0, "a6a5")
        expected = apply_move(after_player, ai_move)
        self.service, *_ = make_service(
            [after_player, expected],
            [ai_move],
            session_store=self.sessions,
            record_writer=self.records,
            config=replace(CONFIG, safety=replace(CONFIG.safety, hardware_enabled=False)),
        )

    def tearDown(self):
        self.service.close()
        self.temporary.cleanup()

    def test_each_game_gets_detailed_workbook_without_blocking_robot_flow(self):
        writing = threading.Event()
        release = threading.Event()
        original_write = self.records._write

        def blocked_write(payload):
            writing.set()
            if not release.wait(timeout=30):
                raise TimeoutError("test did not release workbook writer")
            return original_write(payload)

        # Prove independence from workbook I/O, rather than racing Excel's
        # formatting and disk work against an arbitrary turn deadline.
        with patch.object(self.records, "_write", side_effect=blocked_write):
            try:
                self.service.submit_start("player")
                wait_for(self.service, Phase.PLAYER_TURN.value)
                self.assertTrue(writing.wait(timeout=5))
                self.service.submit_player_done()
                state = wait_for(self.service, Phase.PLAYER_TURN.value)
                self.assertFalse(release.is_set())
            finally:
                release.set()
            self.assertTrue(self.records.flush())

        game_id = state["game"]["id"]
        path = self.records.path_for(game_id)
        self.assertTrue(path.is_file())
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            self.assertEqual(
                workbook.sheetnames,
                ["對局摘要", "著法紀錄", "拍攝紀錄", "Robot動作", "流程紀錄", "棋盤履歷", "引擎紀錄", "辨識盤面"],
            )
            moves = list(workbook["著法紀錄"].iter_rows(min_row=4, values_only=True))
            self.assertEqual([row[5] for row in moves], ["a3a4", "a6a5"])
            self.assertEqual([row[3] for row in moves], ["玩家", "AI"])
            self.assertTrue(all(row[20] and row[21] for row in moves))
            self.assertEqual([row[1] for row in moves], [1, 1])
            captures = list(workbook["拍攝紀錄"].iter_rows(min_row=4, values_only=True))
            self.assertEqual(len(captures), 1)
            self.assertEqual([row[4] for row in captures], ["玩家走子"])
            motions = list(workbook["Robot動作"].iter_rows(min_row=4, values_only=True))
            self.assertEqual(len(motions), 1)
            self.assertEqual(motions[0][6], "a6a5")
            events = list(workbook["流程紀錄"].iter_rows(min_row=4, values_only=True))
            self.assertGreaterEqual(len(events), 3)
            summary_values = {
                row[0]: row[1]
                for row in workbook["對局摘要"].iter_rows(min_row=5, max_col=2, values_only=True)
                if row[0]
            }
            self.assertEqual(summary_values["對局編號"], game_id)
            self.assertEqual(summary_values["已提交著法"], 2)
            self.assertEqual(summary_values["執行模式"], "未啟用")

            saved = json.loads(
                (self.sessions.game_directory(game_id) / "game.json").read_text(encoding="utf-8")
            )
            self.assertEqual(saved["schema_version"], 3)
            self.assertEqual(len(saved["move_history"]), 2)
            self.assertEqual(len(saved["capture_history"]), 1)
            self.assertEqual(len(saved["motion_history"]), 1)
            self.assertTrue(saved["move_history"][0]["fen_before"])
            self.assertTrue(saved["move_history"][0]["fen_after"])
        finally:
            workbook.close()

    def test_monitor_downloads_current_game_workbook(self):
        self.service.submit_start("player")
        state = wait_for(self.service, Phase.PLAYER_TURN.value)
        self.assertTrue(self.records.flush())
        app = create_ui_app(self.service)

        response = app.test_client().get("/api/game/record")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(state["game"]["id"], response.headers["Content-Disposition"])
        response.close()

    def test_new_game_creates_a_separate_workbook(self):
        self.service.submit_start("player")
        first = wait_for(self.service, Phase.PLAYER_TURN.value)["game"]["id"]
        self.service.submit_stop()
        wait_for(self.service, Phase.STOPPED.value)

        self.service.submit_start("player")
        second = wait_for(self.service, Phase.PLAYER_TURN.value)["game"]["id"]
        self.assertNotEqual(first, second)
        self.assertTrue(self.records.flush())

        files = {path.stem for path in self.records.root.rglob("*.xlsx")}
        self.assertEqual(files, {first, second})


if __name__ == "__main__":
    unittest.main()
