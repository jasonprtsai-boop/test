"""Regression paths for one battle spanning detection failures and baselines."""
import json
import tempfile
import unittest
from pathlib import Path

from tmchess.board import Move, apply_move
from tmchess.config import SITE_ROOT
from tmchess.records import GameRecordWriter
from tmchess.service import CommandConflict
from tmchess.session import GameSessionStore
from tmchess.state import Phase
from tmchess.vision import BoardObservationError, CaptureResult
from tests.fakes import make_service, starting_board, wait_for


class BattleContinuityTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(dir=SITE_ROOT / 'runtime')
        self.sessions = GameSessionStore(Path(self.tmp.name) / 'games')
        self.service = None
        self.red = Move((0, 3), (0, 4), 0, 'a3a4')
        self.black = Move((0, 6), (0, 5), 0, 'a6a5')
        self.after_red = apply_move(starting_board(), self.red)
        self.after_black = apply_move(self.after_red, self.black)

    def tearDown(self):
        if self.service: self.service.close()
        self.tmp.cleanup()

    def create(self, outcomes, **kwargs):
        self.service, *rest = make_service(outcomes, session_store=self.sessions, **kwargs)
        return rest

    def saved(self):
        gid = self.service.get_state()['game']['id']
        return json.loads((self.sessions.game_directory(gid) / 'game.json').read_text(encoding='utf-8'))

    def fail_player_capture(self):
        self.service.submit_start('player')
        wait_for(self.service, Phase.PLAYER_TURN.value)
        gid = self.service.get_state()['game']['id']
        self.service.submit_player_done()
        wait_for(self.service, Phase.ERROR.value)
        return gid

    def test_current_board_engine_mode_waits_for_confirmation_for_either_side(self):
        for actor, side, move in (("player", "w", self.red), ("ai", "b", self.black)):
            with self.subTest(side=side):
                robot, _, engine = self.create([starting_board()], suggestions=[move])
                self.service.submit_start_current("red", actor, continuation="new", mode="engine_vs_engine")
                state = wait_for(self.service, Phase.ENGINE_READY.value)
                self.assertEqual(state["side_to_move"], side)
                self.assertEqual(state["controllers"], {"w": "engine", "b": "engine"})
                self.assertEqual(state["execution_mode"], "step")
                self.assertEqual(engine.contexts[-1][1], side)
                self.assertEqual(robot.moves, [])
                self.assertEqual(self.saved()["initial_side_to_move"], side)
                self.service.close()
                self.service = None

    def test_current_board_analysis_black_turn_stays_human_after_rebaseline(self):
        robot, _, engine = self.create([starting_board(), starting_board()])
        self.service.submit_start_current("red", "ai", continuation="new", mode="analysis")
        state = wait_for(self.service, Phase.PLAYER_TURN.value)
        self.assertEqual(state["side_to_move"], "b")
        self.assertEqual(state["controllers"], {"w": "human", "b": "human"})
        self.assertEqual(engine.contexts[-1][1], "b")
        self.service.submit_stop()
        wait_for(self.service, Phase.STOPPED.value)
        self.service.submit_start_current("red", "player", continuation="rebaseline")
        state = wait_for(self.service, Phase.PLAYER_TURN.value)
        self.assertEqual(state["game_mode"], "analysis")
        self.assertEqual(state["execution_mode"], "analysis")
        self.assertEqual(engine.suggest_calls, 0)
        self.assertEqual(robot.moves, [])

    def test_auto_recovery_keeps_saved_analysis_mode_and_black_turn(self):
        self.create([starting_board(), starting_board()])
        self.service.submit_start_current("red", "ai", continuation="new", mode="analysis")
        state = wait_for(self.service, Phase.PLAYER_TURN.value)
        game_id = state["game"]["id"]
        self.service.submit_stop()
        wait_for(self.service, Phase.STOPPED.value)
        self.service.submit_start_current("red", "player", mode="human_red")
        state = wait_for(self.service, Phase.PLAYER_TURN.value)
        self.assertEqual(state["game"]["id"], game_id)
        self.assertEqual(state["game_mode"], "analysis")
        self.assertEqual(state["side_to_move"], "b")

    def test_custom_record_name_survives_rebaseline_and_restart(self):
        writer = GameRecordWriter(Path(self.tmp.name) / 'records', games_root=self.sessions.root)
        self.create([TimeoutError('bad image'), self.after_red], record_writer=writer)
        self.service.submit_start_mode('human_red', record_name='王同學/練習賽')
        wait_for(self.service, Phase.PLAYER_TURN.value)
        gid = self.service.get_state()['game']['id']
        original = writer.path_for(gid)
        self.assertTrue(original.name.startswith('王同學_練習賽_'))
        from datetime import datetime, timedelta
        day = (datetime.strptime(gid.split('-')[0], '%Y%m%dT%H%M%SZ') + timedelta(hours=8)).strftime('%Y-%m-%d')
        self.assertEqual(original.parent, writer.root / '整場對戰' / day)
        self.service.submit_player_done()
        wait_for(self.service, Phase.ERROR.value)
        self.service.submit_start_current('red', 'player', continuation='rebaseline', record_name='不可更名')
        wait_for(self.service, Phase.PLAYER_TURN.value)
        payload = self.saved()
        self.assertEqual(payload['record_name'], '王同學/練習賽')
        self.assertEqual(writer.path_for(payload['game_id']), original)
        self.assertEqual(self.service.get_state()['game']['record_path'], str(original))
        self.service.close()
        self.service = None
        restarted = GameRecordWriter(writer.root, games_root=self.sessions.root)
        try:
            restarted.schedule(payload)
            self.assertEqual(restarted.path_for(payload['game_id']), original)
            self.assertTrue(restarted.flush())
            self.assertTrue(original.is_file())
        finally:
            restarted.close()

    def test_current_board_after_detection_failure_recovers_same_game_and_move(self):
        self.create([TimeoutError('V2 missing'), self.after_red, self.after_black], suggestions=[self.black])
        gid = self.fail_player_capture()
        self.service.submit_start_current('red', 'player')
        wait_for(self.service, Phase.PLAYER_TURN.value)
        p = self.saved()
        self.assertEqual(p['game_id'], gid)
        self.assertEqual([m['uci'] for m in p['move_history']], ['a3a4', 'a6a5'])
        self.assertEqual(len(p['capture_history']), 2)
        self.assertEqual(len(list(self.sessions.root.iterdir())), 1)

    def test_explicit_rebaseline_retains_battle_identity_and_clears_old_image(self):
        self.create([TimeoutError('bad image'), self.after_red])
        gid = self.fail_player_capture()
        self.service.submit_start_current('red', 'player', continuation='rebaseline')
        wait_for(self.service, Phase.PLAYER_TURN.value)
        p = self.saved()
        self.assertNotEqual(p['game_id'], gid)
        self.assertEqual(p['battle_id'], gid)
        self.assertEqual(p['parent_game_id'], gid)
        self.assertEqual(p['segment_index'], 2)
        self.assertTrue(p['history_gap'])
        self.assertIsNone(p['event_history'][0]['capture_id'])
        self.assertEqual(p['capture_history'][0]['accepted_board'],self.after_red)
        self.assertEqual(p['capture_history'][0]['acceptance_status'],'accepted_initial_baseline')

    def test_explicit_new_battle_does_not_link_to_failed_game(self):
        self.create([TimeoutError('bad image'), starting_board()])
        gid = self.fail_player_capture()
        self.service.submit_start_current('red', 'player', continuation='new')
        wait_for(self.service, Phase.PLAYER_TURN.value)
        p = self.saved()
        self.assertNotEqual(p['battle_id'], gid)
        self.assertEqual(p['battle_id'], p['game_id'])
        self.assertIsNone(p['parent_game_id'])

    def test_baseline_validation_failure_preserves_observation_and_metrics(self):
        bad = starting_board(); bad[4][4] = 9
        result = CaptureResult('sample', bad, Path('v1.jpg'), Path('v2.jpg'), [[.8]*9 for _ in range(10)],
                               capture_seconds=1.2, processing_seconds=.3, total_seconds=1.5)
        self.create([BoardObservationError('too many red rooks', result)])
        self.service.submit_start_current('red', 'player')
        wait_for(self.service, Phase.ERROR.value)
        c = self.saved()['capture_history'][0]
        self.assertEqual(c['observed_board'], bad)
        self.assertEqual(c['processing_seconds'], .3)
        self.assertEqual(c['total_seconds'], 1.5)
        self.assertIsNotNone(c['workflow_seconds'])
        self.assertEqual(c['detected_pieces'], 33)
        self.assertEqual(c['confidence_grid'][4][4], .8)
        self.assertEqual(c['status'], 'failed')

    def test_timeout_reports_existing_artifacts_instead_of_fictitious_paths(self):
        self.create([TimeoutError('V2 missing')])
        self.fail_player_capture()
        c = self.saved()['capture_history'][0]
        self.assertIsNone(c['result_json'])
        self.assertIsNone(c['v1_raw'])
        self.assertEqual(c['artifact_status']['v2_raw'], 'not_received')
        self.assertIsNotNone(c['total_seconds'])

    def test_engine_failure_is_recorded_even_without_committed_move(self):
        self.create([self.after_red], engine_error='suggest')
        self.service.submit_start('player'); wait_for(self.service, Phase.PLAYER_TURN.value)
        self.service.submit_player_done(); wait_for(self.service, Phase.ERROR.value)
        q = self.saved()['engine_history'][-1]
        self.assertEqual(q['operation'], 'suggest')
        self.assertEqual(q['status'], 'failed')
        self.assertEqual(q['error_type'], 'RuntimeError')
        self.assertIsNotNone(q['duration_seconds'])

    def test_ai_legality_and_elapsed_time_and_engine_queries_are_saved(self):
        self.create([self.after_red, self.after_black], suggestions=[self.black])
        self.service.submit_start('player'); wait_for(self.service, Phase.PLAYER_TURN.value)
        self.service.submit_player_done(); wait_for(self.service, Phase.PLAYER_TURN.value)
        p = self.saved(); ai = p['move_history'][1]
        self.assertIsNotNone(ai['legality_seconds'])
        self.assertIsNotNone(ai['turn_seconds'])
        self.assertTrue(ai['move_attempt_id'])
        self.assertTrue(p['engine_history'])
        self.assertEqual(p['motion_history'][0]['verification_status'], 'skipped')
        self.assertEqual(p['motion_history'][0]['committed_move_sequence'], 2)

    def test_excel_has_engine_history_and_explicit_missing_metric_labels(self):
        p = dict(game_id='20261020T010101Z-12345678', phase='ERROR', move_history=[
            dict(actor='player', uci='a3a4', source=[0,3], destination=[0,4]),
            dict(actor='ai', uci='a6a5', source=[0,6], destination=[0,5])],
            engine_history=[dict(sequence=1, operation='suggest', status='failed', error_message='engine died')])
        w = GameRecordWriter._build_workbook(p)
        try:
            self.assertIn('引擎紀錄', w.sheetnames)
            self.assertEqual(w['著法紀錄'].cell(4,30).value, '不適用')
            self.assertEqual(w['著法紀錄'].cell(5,30).value, '未記錄')
            self.assertEqual(w['引擎紀錄'].cell(4,5).value, 'failed')
        finally: w.close()

    def test_continuation_survives_process_restart(self):
        self.create([TimeoutError('bad image'), self.after_red])
        gid = self.fail_player_capture()
        self.service.submit_start_current('red', 'player', continuation='rebaseline')
        wait_for(self.service, Phase.PLAYER_TURN.value)
        segment = self.saved()['game_id']
        self.service.close()
        self.create([self.after_red])
        self.service.submit_start_current('red', 'player')
        wait_for(self.service, Phase.PLAYER_TURN.value)
        self.assertEqual(self.saved()['battle_id'], gid)
        self.assertEqual(self.saved()['game_id'], segment)
        self.assertEqual(self.saved()['segment_index'], 2)

    def test_auto_recovery_rejects_color_change_without_creating_game(self):
        self.create([TimeoutError('bad image')])
        gid = self.fail_player_capture()
        with self.assertRaises(CommandConflict):
            self.service.submit_start_current('black', 'player')
        self.assertEqual(self.saved()['game_id'], gid)
        self.assertFalse(self.service._busy)

    def test_completed_move_resumes_without_replaying_robot(self):
        robot, vision, _ = self.create([self.after_red, self.after_black], suggestions=[self.black])
        self.service.submit_start('player'); wait_for(self.service, Phase.PLAYER_TURN.value)
        self.service.submit_player_done(); wait_for(self.service, Phase.PLAYER_TURN.value)
        self.service.submit_stop(); wait_for(self.service, Phase.STOPPED.value)
        gid = self.saved()['game_id']
        self.service.submit_start_current('red', 'player')
        wait_for(self.service, Phase.PLAYER_TURN.value)
        self.assertEqual(self.saved()['game_id'], gid)
        self.assertEqual(len(robot.moves), 1)
        self.assertEqual(len(self.saved()['move_history']), 2)

    def test_partial_photo_survives_timeout_and_is_listed(self):
        _, vision, _ = self.create([])
        def partial(capture_id, trigger):
            trigger()
            root = self.sessions.game_directory(vision.selected_game)
            (root / f'{capture_id}_v1.jpg').write_bytes(b'partial-image')
            raise TimeoutError('V2 missing')
        vision.capture = partial
        self.fail_player_capture()
        c = self.saved()['capture_history'][0]
        self.assertEqual(c['artifact_status']['v1_raw'], 'available')
        self.assertIsNotNone(c['v1_raw'])
        self.assertIsNone(c['v2_raw'])

    def test_aggregate_workbook_keeps_all_segments_and_no_gap_is_invented_as_moves(self):
        from openpyxl import load_workbook
        self.create([TimeoutError('bad image'), self.after_red])
        gid = self.fail_player_capture()
        self.service.submit_start_current('red', 'player', continuation='rebaseline')
        wait_for(self.service, Phase.PLAYER_TURN.value)
        p = self.saved()
        writer = GameRecordWriter(Path(self.tmp.name)/'records', games_root=self.sessions.root)
        try:
            writer.schedule(p)
            self.assertTrue(writer.flush())
            path = writer.path_for(p['game_id'])
            self.assertTrue(path.name.startswith('象棋對戰_'))
            self.assertTrue(path.name.endswith(gid.split('-')[-1] + '.xlsx'))
            w = load_workbook(path, read_only=True)
            try:
                self.assertEqual(w['對戰片段'].max_row, 5)
                self.assertEqual(w['拍攝紀錄'].max_row, 5)
                self.assertEqual(w['著法紀錄'].max_row, 3)
                self.assertIn('棋盤履歷1',w.sheetnames)
                self.assertIn('棋盤履歷2',w.sheetnames)
            finally: w.close()
        finally: writer.close()

    def test_aggregate_rejects_cross_battle_parent(self):
        writer = GameRecordWriter(Path(self.tmp.name)/'records', games_root=self.sessions.root)
        first = '20261020T010101Z-12345678'; second = '20261020T010102Z-12345678'
        self.sessions.create(first)
        self.sessions.save(first, dict(game_id=first,battle_id=first))
        try:
            with self.assertRaisesRegex(ValueError, '來源不一致'):
                writer._battle_segments(dict(game_id=second,battle_id=second,parent_game_id=first))
        finally: writer.close()

    def test_older_segment_export_cannot_replace_newer_battle_head(self):
        self.create([TimeoutError('bad image'),self.after_red])
        gid=self.fail_player_capture(); first=self.saved()
        self.service.submit_start_current('red','player',continuation='rebaseline')
        wait_for(self.service,Phase.PLAYER_TURN.value); second=self.saved()
        writer=GameRecordWriter(Path(self.tmp.name)/'records',games_root=self.sessions.root)
        try:
            writer.schedule(second); self.assertTrue(writer.flush())
            writer.schedule(first); self.assertTrue(writer.flush())
            from openpyxl import load_workbook
            w=load_workbook(writer.path_for(gid),read_only=True)
            try: self.assertEqual(w['對戰片段'].max_row,5)
            finally:w.close()
        finally:writer.close()

    def test_operator_confirmation_is_preserved_after_restart(self):
        self.create([])
        self.service.submit_start('player'); wait_for(self.service,Phase.PLAYER_TURN.value)
        p=self.saved();self.service.close()
        p['motion_state_uncertain']=True;p['phase']='STOPPED'
        self.sessions.save(p['game_id'],p)
        self.create([])
        self.service.submit_confirm_safe_stop(True)
        saved=json.loads((self.sessions.game_directory(p['game_id'])/'game.json').read_text(encoding='utf-8'))
        self.assertFalse(saved['motion_state_uncertain'])
        self.assertTrue(any(e.get('event_type')=='operator_safe_stop_confirmation' for e in saved['event_history']))

    def test_report_rebuild_derives_metadata_without_mutating_source(self):
        from tools.rebuild_record_reports import enrich
        initial=starting_board()
        source=dict(game_id='20261020T010101Z-12345678',initial_board=initial,initial_side_to_move='w',
                    board=self.after_red,side_to_move='w',move_history=[dict(uci='a3a4',side='b')])
        before=json.dumps(source)
        p,audit=enrich(source,Path(self.tmp.name),correct_metadata=True)
        self.assertEqual(p['side_to_move'],'b')
        self.assertEqual(p['move_history'][0]['side'],'w')
        self.assertTrue(audit['corrections'])
        self.assertEqual(json.dumps(source),before)


if __name__ == '__main__': unittest.main()
